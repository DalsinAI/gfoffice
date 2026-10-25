#!/usr/bin/env python3
"""What a sheet holds beside its cells, read and kept on saving:
test_sheet_features.py BOOK_DUMP OWCONVERT

On fixtures/features-openpyxl.xlsx (make_feature_fixtures.py: a workbook
made up to have what the lab's yardstick has) and on the .xlsx OWConvert
writes from it: book_dump's lines must say the same of both, the file
written must be what openpyxl (when installed) and LibreOffice (when
installed; it is not needed) read back with the same, and ODS must say
what it leaves out.

  - conditional formatting: cellIs, expression and colour scale rules
    with their looks (the styles' dxfs), a data bar and a "contains text"
    rule kept as the file wrote them, priorities and stopIfTrue;
  - comments: their text and author, written with the boxes Excel shows
    them in (a VML drawing);
  - data validation: lists, numbers, dates, text lengths and formulas,
    their operators, messages and error styles;
  - the AutoFilter: its range, and what it filters and sorts by;
  - pictures and charts: a sheet's drawing and every part it leads to
    (charts, their styles and colours, pictures) kept as they were and
    written back the same, with the content types and relationships that
    reach them; a drawing whose parts are missing left out whole;
  - charts: what owf_sheet_charts reads of them (where, kind, title,
    legend, axis, and each series' cells, colour and cached values).
MIT, Copyright (c) 2026 AmigaChrome Team.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.dom.minidom

HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, "fixtures", "features-openpyxl.xlsx")
CHARTS = os.path.join(HERE, "fixtures", "charts-openpyxl.xlsx")
DUMP = os.path.abspath(sys.argv[1])
OWCONVERT = os.path.abspath(sys.argv[2])
SOFFICE = shutil.which("soffice") if os.environ.get("OWF_NO_SOFFICE") is None else None
failures = []
checks = 0


def check(name, condition, detail=""):
    global checks
    checks += 1
    if not condition:
        failures.append(name)
        print("FAIL %s %s" % (name, str(detail)[:2000]))


def run(args, timeout=300):
    p = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=timeout)
    return p.returncode, p.stdout.decode("utf-8", "replace"), p.stderr.decode("utf-8", "replace")


def dump(path):
    rc, out, err = run([DUMP, path])
    check("book_dump " + os.path.basename(path), rc == 0, out + err)
    return out


def lines(text, kind):
    return [l for l in text.splitlines() if l.startswith(kind + " ")]


def sheet_lines(text, sheet, kind):
    """book_dump's lines of one kind under sheet number `sheet`."""
    out, current = [], -1
    for l in text.splitlines():
        if l.startswith("sheet "):
            current = int(l.split()[1])
        elif current == sheet and l.startswith(kind + " "):
            out.append(l)
    return out


def convert(src, dst):
    rc, out, err = run([OWCONVERT, "--report", src, dst])
    check("owconvert %s -> %s" % (os.path.basename(src), os.path.basename(dst)), rc == 0, out + err)
    return out + err


def well_formed(path):
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    xml.dom.minidom.parseString(z.read(name))
                except Exception as e:  # noqa: BLE001 - any parse error is a failure
                    check("well-formed " + name, False, e)


# ---- Conditional formatting ----

def test_condfmt(work):
    src = dump(FIX)
    cf = sheet_lines(src, 0, "cf")
    check("cf: ten rules read", len(cf) == 10, "\n".join(cf))
    by_priority = {}
    for l in cf:
        m = re.search(r"priority=(\d+)", l)
        by_priority[int(m.group(1))] = l
    # cellIs greater than 800 on F7:F242: dark red text on light red
    l = by_priority.get(1, "")
    check("cf 1 cellIs greaterThan", "type=0 op=4" in l and "ranges=6,5-241,5" in l and 'formula "800"' in l, l)
    check("cf 1 look", "colour=9C0006 fill=FFC7CE" in l, l)
    l = by_priority.get(2, "")
    check("cf 2 between", "type=0 op=0" in l and 'formula "100" formula "200"' in l and "fill=C6EFCE" in l, l)
    l = by_priority.get(3, "")
    check("cf 3 expression", 'type=1' in l and 'formula "$D7=\\"No\\""' in l and "ranges=6,0-241,8" in l, l)
    check("cf 3 look: italic grey", "set=2 flags=2 colour=808080" in l, l)
    l = by_priority.get(4, "")
    check("cf 4 three-colour scale", "type=2" in l and 'point 0 "" F8696B point 4 "50" FFEB84 point 1 "" 63BE7B' in l, l)
    l = by_priority.get(5, "")
    check("cf 5 two-colour scale", "type=2" in l and 'point 0 "" FFFFFF point 1 "" 5A8AC6' in l, l)
    l = by_priority.get(6, "")
    check("cf 6 data bar kept", "type=3" in l and 'xml "<cfRule type=\\"dataBar\\"><dataBar>' in l and "dxfId" not in l, l)
    l = by_priority.get(7, "")
    check("cf 7 contains text kept, shown by its formula",
          "type=3" in l and 'formula "NOT(ISERROR(SEARCH(\\"Review\\",I7)))"' in l and "containsText" in l, l)
    check("cf 7 look", "set=1 flags=1 colour=9C0006 fill=FFC7CE" in l, l)
    l = by_priority.get(8, "")
    check("cf 8 stopIfTrue", "type=0 op=2" in l and "stop=1" in l and 'formula "\\"To confirm\\""' in l, l)
    rc, out, err = run([DUMP, FIX])
    check("cf: no 'not read yet' note", "Conditional formatting is not read" not in out, out)
    check("cf: kept-not-shown note", "kept but not shown" in out, out)

    # saved as .xlsx and read again: the same rules
    xlsx = os.path.join(work, "features.xlsx")
    convert(FIX, xlsx)
    well_formed(xlsx)
    back = dump(xlsx)
    check("cf: the same after saving", sheet_lines(back, 0, "cf") == cf,
          "\n".join(sheet_lines(back, 0, "cf")) + "\n---\n" + "\n".join(cf))
    with zipfile.ZipFile(xlsx) as z:
        styles = z.read("xl/styles.xml").decode()
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
    m = re.search(r'<dxfs count="(\d+)"', styles)
    check("cf: one dxf a rule with a look", m and m.group(1) == "7", styles[-1500:])
    check("cf: conditionalFormatting after mergeCells", sheet.find("</mergeCells>") < sheet.find("<conditionalFormatting")
          < sheet.find("<pageMargins"), "")
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl:
        ws = openpyxl.load_workbook(xlsx)["Systems"]
        rules = sorted([(str(cfr.sqref), rule.type, rule.operator, tuple(rule.formula or ()), bool(rule.stopIfTrue))
                        for cfr in ws.conditional_formatting for rule in cfr.rules], key=str)
        ws0 = openpyxl.load_workbook(FIX)["Systems"]
        want = sorted([(str(cfr.sqref), rule.type, rule.operator, tuple(rule.formula or ()), bool(rule.stopIfTrue))
                       for cfr in ws0.conditional_formatting for rule in cfr.rules], key=str)
        check("cf: openpyxl reads the same rules back", rules == want, "%s\n---\n%s" % (rules, want))
    if SOFFICE:
        def lo_formats(path, outdir):
            """LibreOffice's ODS of a file: its conditional formats' parts, counted by kind."""
            profile = tempfile.mkdtemp(prefix="owf-lo-")
            try:
                os.makedirs(outdir, exist_ok=True)
                subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", "ods",
                                "--outdir", outdir, path], check=True, stdout=subprocess.DEVNULL,
                               stderr=subprocess.DEVNULL, timeout=300)
                name = os.path.splitext(os.path.basename(path))[0] + ".ods"
                with zipfile.ZipFile(os.path.join(outdir, name)) as z:
                    content = z.read("content.xml").decode()
            finally:
                shutil.rmtree(profile, ignore_errors=True)
            return {k: content.count("<calcext:" + k) for k in ("conditional-format ", "condition ", "color-scale ",
                                                                  "data-bar ", "color-scale-entry ")}
        try:
            mine = lo_formats(xlsx, os.path.join(work, "lo-mine"))
            theirs = lo_formats(FIX, os.path.join(work, "lo-fixture"))
            check("cf: LibreOffice reads the rules written as it reads openpyxl's", mine == theirs and
                  mine["conditional-format "] == 8 and mine["condition "] == 7 and mine["data-bar "] == 1, "%s\n%s" % (mine, theirs))
        except Exception as e:  # noqa: BLE001
            check("cf: LibreOffice converts the file written", False, e)

    # ODS: said to be left out
    report = convert(FIX, os.path.join(work, "features.ods"))
    check("cf: ODS says it leaves them out", "Conditional formatting is kept in Excel workbooks" in report, report)


# ---- Comments ----

def test_comments(work):
    src = dump(FIX)
    cms = sheet_lines(src, 0, "comment")
    want = ['comment 6,0 author "Team 5" text "Main identity service.\\nOwner to confirm."',
            'comment 8,2 author "Team 5" text "Hosted by the supplier"',
            'comment 11,5 author "Finance" text "Cost per year, before VAT"']
    check("comments read", cms == want, "\n".join(cms))
    rc, out, err = run([DUMP, FIX])
    check("comments: no 'not read yet' note", "Comments on cells are not read" not in out, out)
    xlsx = os.path.join(work, "comments.xlsx")
    convert(FIX, xlsx)
    well_formed(xlsx)
    check("comments: the same after saving", sheet_lines(dump(xlsx), 0, "comment") == want, "")
    with zipfile.ZipFile(xlsx) as z:
        names = z.namelist()
        types = z.read("[Content_Types].xml").decode()
        rels = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode() if "xl/worksheets/_rels/sheet1.xml.rels" in names else ""
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        vml = z.read("xl/drawings/vmlDrawing1.vml").decode() if "xl/drawings/vmlDrawing1.vml" in names else ""
    check("comments: the parts", "xl/comments1.xml" in names and "/xl/comments1.xml" in types and 'Extension="vml"' in types, names)
    check("comments: the sheet's relationships", "comments1.xml" in rels and "vmlDrawing1.vml" in rels, rels)
    check("comments: the legacy drawing, last", sheet.rstrip().endswith('<legacyDrawing r:id="rId2"/></worksheet>'), sheet[-200:])
    check("comments: a note shape each", vml.count('ObjectType="Note"') == 3 and "<x:Row>6</x:Row><x:Column>0</x:Column>" in vml, vml[:500])
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl:
        ws = openpyxl.load_workbook(xlsx)["Systems"]
        got = [(c, ws[c].comment.text if ws[c].comment else None, ws[c].comment.author if ws[c].comment else None)
               for c in ("A7", "C9", "F12", "B7")]
        check("comments: openpyxl reads them back", got == [("A7", "Main identity service.\nOwner to confirm.", "Team 5"),
                                                           ("C9", "Hosted by the supplier", "Team 5"),
                                                           ("F12", "Cost per year, before VAT", "Finance"), ("B7", None, None)], got)
    if SOFFICE:
        profile = tempfile.mkdtemp(prefix="owf-lo-")
        try:
            outdir = os.path.join(work, "lo-comments")
            os.makedirs(outdir, exist_ok=True)
            subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", "ods",
                            "--outdir", outdir, xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=300)
            with zipfile.ZipFile(os.path.join(outdir, "comments.ods")) as z:
                content = z.read("content.xml").decode()
            check("comments: LibreOffice reads the three", content.count("<office:annotation") == 3 and
                  "Cost per year, before VAT" in content, content.count("<office:annotation"))
        except Exception as e:  # noqa: BLE001
            check("comments: LibreOffice converts the file written", False, e)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
    report = convert(FIX, os.path.join(work, "comments.ods"))
    check("comments: ODS says it leaves them out", "Comments on cells are kept in Excel workbooks" in report, report)


# ---- Data validation ----

VALIDATIONS = [
    'validation type=3 op=0 blank=1 nodropdown=0 prompt=0 error=0 style=0 ranges=6,3-241,3 formula "\\"Yes,No\\"" '
    'prompt "Hosted" "Hosted by us?" message "Not a choice" "Yes or No"',
    'validation type=2 op=0 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=6,5-241,5 formula "0" formula "100000"',
    'validation type=4 op=4 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=6,6-241,6 formula "DATE(2020,1,1)"',
    'validation type=6 op=7 blank=0 nodropdown=0 prompt=0 error=1 style=1 ranges=6,8-241,8 formula "200"',
    'validation type=3 op=0 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=6,4-241,4 formula "Lists!$A$1:$A$5"',
    'validation type=1 op=0 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=6,7-19,7;29,7-39,7 formula "0" formula "10"',
    'validation type=7 op=0 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=6,0-241,0 formula "LEN(A7)>0"',
]


def openpyxl_validations(path):
    import openpyxl
    ws = openpyxl.load_workbook(path)["Systems"]
    return sorted((str(v.sqref), v.type, v.operator or "between", v.formula1, v.formula2, bool(v.allow_blank),
                   v.promptTitle, v.prompt, v.errorTitle, v.error, v.errorStyle or "stop", bool(v.showErrorMessage))
                  for v in ws.data_validations.dataValidation)


def test_validation(work):
    src = dump(FIX)
    got = sheet_lines(src, 0, "validation")
    check("validation: the seven read", got == VALIDATIONS, "\n".join(got))
    rc, out, err = run([DUMP, FIX])
    check("validation: no 'not read yet' note", "Data validation (lists to pick from, limits) is not read" not in out, out)
    xlsx = os.path.join(work, "validation.xlsx")
    convert(FIX, xlsx)
    well_formed(xlsx)
    check("validation: the same after saving", sheet_lines(dump(xlsx), 0, "validation") == VALIDATIONS, "")
    with zipfile.ZipFile(xlsx) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
    check("validation: after the conditional formats, before the margins",
          sheet.rfind("</conditionalFormatting>") < sheet.find('<dataValidations count="7">') < sheet.find("<pageMargins"), "")
    try:
        import openpyxl  # noqa: F401
        have = True
    except ImportError:
        have = False
    if have:
        a, b = openpyxl_validations(xlsx), openpyxl_validations(FIX)
        check("validation: openpyxl reads them back", a == b, "%s\n---\n%s" % (a, b))
    if SOFFICE:
        profile = tempfile.mkdtemp(prefix="owf-lo-")
        try:
            outdir = os.path.join(work, "lo-validation")
            os.makedirs(outdir, exist_ok=True)
            subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", "ods",
                            "--outdir", outdir, xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=300)
            with zipfile.ZipFile(os.path.join(outdir, "validation.ods")) as z:
                content = z.read("content.xml").decode()
            n = content.count("<table:content-validation ")
            check("validation: LibreOffice reads the seven", n == 7 and "Hosted by us?" in content, n)
        except Exception as e:  # noqa: BLE001
            check("validation: LibreOffice converts the file written", False, e)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
    report = convert(FIX, os.path.join(work, "validation.ods"))
    check("validation: ODS says it leaves them out", "Data validation is kept in Excel workbooks" in report, report)


# ---- The AutoFilter ----

AF_KEPT = ('<filterColumn colId="3"><filters><filter val="Yes"/></filters></filterColumn>'
           '<sortState ref="A7:I242"><sortCondition ref="A7:A242"/></sortState>')


def test_autofilter(work):
    src = dump(FIX)
    check("autofilter read", sheet_lines(src, 0, "autofilter") == ['autofilter 5,0-241,8 ""'], sheet_lines(src, 0, "autofilter"))
    rc, out, err = run([DUMP, FIX])
    check("autofilter: no 'not read yet' note", "AutoFilter buttons are not read" not in out, out)
    # the fixture with a filter and a sort in it, as Excel writes them
    with zipfile.ZipFile(FIX) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(r'<autoFilter ref="A6:I242"\s*/>', '<autoFilter ref="A6:I242">' + AF_KEPT + '<extLst><ext uri="x"/></extLst></autoFilter>', sheet)
    check("autofilter: the fixture made to filter", AF_KEPT in sheet, sheet[:300])
    filtered = os.path.join(work, "filtered.xlsx")
    with zipfile.ZipFile(filtered, "w") as z:
        for n, data in parts.items():
            z.writestr(n, sheet.encode() if n == "xl/worksheets/sheet1.xml" else data)
    want = ['autofilter 5,0-241,8 "' + AF_KEPT.replace('"', '\\"') + '"']
    check("autofilter: what it filters by kept (its extensions left out)", sheet_lines(dump(filtered), 0, "autofilter") == want,
          sheet_lines(dump(filtered), 0, "autofilter"))
    xlsx = os.path.join(work, "autofilter.xlsx")
    convert(filtered, xlsx)
    well_formed(xlsx)
    check("autofilter: the same after saving", sheet_lines(dump(xlsx), 0, "autofilter") == want, sheet_lines(dump(xlsx), 0, "autofilter"))
    with zipfile.ZipFile(xlsx) as z:
        written = z.read("xl/worksheets/sheet1.xml").decode()
    check("autofilter: after the cells, before the merged ones",
          written.find("</sheetData>") < written.find('<autoFilter ref="A6:I242">') < written.find("<mergeCells"), "")
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl:
        ws = openpyxl.load_workbook(xlsx)["Systems"]
        af = ws.auto_filter
        check("autofilter: openpyxl reads it back", af.ref == "A6:I242" and len(af.filterColumn) == 1 and
              af.filterColumn[0].colId == 3 and af.filterColumn[0].filters.filter == ["Yes"], (af.ref, af.filterColumn))
    if SOFFICE:
        profile = tempfile.mkdtemp(prefix="owf-lo-")
        try:
            outdir = os.path.join(work, "lo-autofilter")
            os.makedirs(outdir, exist_ok=True)
            subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", "ods",
                            "--outdir", outdir, xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=300)
            with zipfile.ZipFile(os.path.join(outdir, "autofilter.ods")) as z:
                content = z.read("content.xml").decode()
            check("autofilter: LibreOffice reads it", 'table:display-filter-buttons="true"' in content and
                  "Systems.A6:" in content.replace("'", "").replace("$", ""), content[content.find("<table:database-range"):][:300])
        except Exception as e:  # noqa: BLE001
            check("autofilter: LibreOffice converts the file written", False, e)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
    report = convert(FIX, os.path.join(work, "autofilter.ods"))
    check("autofilter: ODS says it leaves it out", "AutoFilters are kept in Excel workbooks" in report, report)


# ---- Pictures and charts ----

DML = "http://schemas.openxmlformats.org/drawingml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def png_1x1():
    import struct
    import zlib

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) +
            chunk(b"IDAT", zlib.compress(b"\x00\xd0\x10\x10")) + chunk(b"IEND", b""))


def picture_anchor(row, rid, link=None):
    click = '<a:hlinkClick xmlns:a="%s" xmlns:r="%s" r:id="%s"/>' % (DML, RNS, link) if link else ""
    return ('<twoCellAnchor editAs="oneCell"><from><col>3</col><colOff>0</colOff><row>%d</row><rowOff>0</rowOff></from>'
            '<to><col>4</col><colOff>0</colOff><row>%d</row><rowOff>0</rowOff></to><pic><nvPicPr>'
            '<cNvPr id="%d" name="Picture %d">%s</cNvPr><cNvPicPr/></nvPicPr><blipFill>'
            '<a:blip xmlns:a="%s" xmlns:r="%s" r:embed="%s"/><a:stretch xmlns:a="%s"><a:fillRect/></a:stretch></blipFill>'
            '<spPr><a:prstGeom xmlns:a="%s" prst="rect"><a:avLst/></a:prstGeom></spPr></pic><clientData/></twoCellAnchor>'
            % (row, row + 2, row + 2, row + 2, click, DML, RNS, rid, DML, DML))


def drawn_book(parts):
    """The fixture as Excel writes a chart (its style and colours parts beside it) with a picture under it, a
    link on the picture, and a second drawing (the same picture) on the sheet with the comments."""
    p = dict(parts)
    p["xl/media/image1.png"] = png_1x1()
    p["xl/charts/style1.xml"] = b'<cs:chartStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" id="201"/>'
    p["xl/charts/colors1.xml"] = (b'<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" '
                                  b'meth="cycle" id="10"/>')
    p["xl/charts/_rels/chart1.xml.rels"] = (
        '<Relationships xmlns="%s"><Relationship Id="rId2" Type="http://schemas.microsoft.com/office/2011/relationships/'
        'chartColorStyle" Target="colors1.xml"/><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/'
        'relationships/chartStyle" Target="style1.xml"/></Relationships>' % RELS).encode()
    drawing = p["xl/drawings/drawing1.xml"].decode()
    p["xl/drawings/drawing1.xml"] = drawing.replace("</wsDr>", picture_anchor(18, "rId2", "rId3") + "</wsDr>").encode()
    p["xl/drawings/_rels/drawing1.xml.rels"] = p["xl/drawings/_rels/drawing1.xml.rels"].replace(b"</Relationships>", (
        '<Relationship Id="rId2" Type="%s/image" Target="../media/image1.png"/>'
        '<Relationship Id="rId3" Type="%s/hyperlink" Target="https://example.org/a?b=1&amp;c=2" TargetMode="External"/>'
        '</Relationships>' % (RNS, RNS)).encode())
    p["xl/drawings/drawing2.xml"] = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
                                     + picture_anchor(1, "rId1") + "</xdr:wsDr>").replace(
        "<twoCellAnchor", "<xdr:twoCellAnchor").replace("</twoCellAnchor>", "</xdr:twoCellAnchor>")
    p["xl/drawings/drawing2.xml"] = re.sub(r"<(/?)(from|to|col|colOff|row|rowOff|pic|nvPicPr|cNvPr|cNvPicPr|blipFill|spPr|clientData)([ >/])",
                                           r"<\1xdr:\2\3", p["xl/drawings/drawing2.xml"]).encode()
    p["xl/drawings/_rels/drawing2.xml.rels"] = ('<Relationships xmlns="%s"><Relationship Id="rId1" Type="%s/image" '
                                                'Target="/xl/media/image1.png"/></Relationships>' % (RELS, RNS)).encode()
    sheet = p["xl/worksheets/sheet1.xml"].decode()
    p["xl/worksheets/sheet1.xml"] = sheet.replace("<legacyDrawing", '<drawing xmlns:r="%s" r:id="pics"/><legacyDrawing' % RNS).encode()
    p["xl/worksheets/_rels/sheet1.xml.rels"] = p["xl/worksheets/_rels/sheet1.xml.rels"].replace(b"</Relationships>", (
        '<Relationship Type="%s/drawing" Target="../drawings/drawing2.xml" Id="pics"/></Relationships>' % RNS).encode())
    types = p["[Content_Types].xml"].decode()
    p["[Content_Types].xml"] = types.replace("</Types>", (
        '<Default Extension="PNG" ContentType="image/png"/>'
        '<Override PartName="/xl/drawings/drawing2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
        '<Override PartName="/xl/charts/style1.xml" ContentType="application/vnd.ms-office.chartstyle+xml"/>'
        '<Override PartName="/xl/charts/colors1.xml" ContentType="application/vnd.ms-office.chartcolorstyle+xml"/></Types>')).encode()
    return p


KEPT = {
    "xl/drawings/drawing2.xml": "application/vnd.openxmlformats-officedocument.drawing+xml",
    "xl/drawings/_rels/drawing2.xml.rels": "application/vnd.openxmlformats-package.relationships+xml",
    "xl/media/image1.png": "image/png",
    "xl/drawings/drawing1.xml": "application/vnd.openxmlformats-officedocument.drawing+xml",
    "xl/drawings/_rels/drawing1.xml.rels": "application/vnd.openxmlformats-package.relationships+xml",
    "xl/charts/chart1.xml": "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
    "xl/charts/_rels/chart1.xml.rels": "application/vnd.openxmlformats-package.relationships+xml",
    "xl/charts/colors1.xml": "application/vnd.ms-office.chartcolorstyle+xml",
    "xl/charts/style1.xml": "application/vnd.ms-office.chartstyle+xml",
}


def write_zip(path, parts):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


def kept_parts(text):
    out = {}
    for l in lines(text, "part"):
        m = re.match(r'part "([^"]*)" "([^"]*)" length=(\d+)$', l)
        check("drawings: part line " + l, m is not None)
        if m:
            out[m.group(1)] = m.group(2)
    return out


def content_types(z):
    types = z.read("[Content_Types].xml").decode()
    return dict(re.findall(r'<Override PartName="/([^"]+)" ContentType="([^"]+)"/>', types))


def test_drawings(work):
    src = dump(FIX)
    check("drawings: the chart's sheet has its drawing", sheet_lines(src, 7, "drawing") == ['drawing "xl/drawings/drawing1.xml"'],
          sheet_lines(src, 7, "drawing"))
    check("drawings: its parts kept", kept_parts(src) == {n: KEPT[n] for n in ("xl/drawings/drawing1.xml",
          "xl/drawings/_rels/drawing1.xml.rels", "xl/charts/chart1.xml")}, kept_parts(src))
    check("drawings: said to be kept, not 'not read yet'", "kept for saving as an Excel workbook" in src and
          "not read yet\" x" not in "\n".join(l for l in lines(src, "note") if "ictures" in l), lines(src, "note"))
    with zipfile.ZipFile(FIX) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    drawn = drawn_book(parts)
    book = os.path.join(work, "drawn.xlsx")
    write_zip(book, drawn)
    got = dump(book)
    check("drawings: each part kept once, with its content type", kept_parts(got) == KEPT, kept_parts(got))
    check("drawings: the sheets' drawings", sheet_lines(got, 0, "drawing") == ['drawing "xl/drawings/drawing2.xml"'] and
          sheet_lines(got, 7, "drawing") == ['drawing "xl/drawings/drawing1.xml"'], (sheet_lines(got, 0, "drawing"),
                                                                                       sheet_lines(got, 7, "drawing")))
    xlsx = os.path.join(work, "drawn-saved.xlsx")
    convert(book, xlsx)
    well_formed(xlsx)
    with zipfile.ZipFile(xlsx) as z:
        names = z.namelist()
        check("drawings: each part written once", len(names) == len(set(names)), sorted(names))
        for n in KEPT:
            check("drawings: written as it was: " + n, n in names and z.read(n) == drawn[n], n)
        types = content_types(z)
        for n, t in KEPT.items():
            if not n.endswith(".rels"):
                check("drawings: content type of " + n, types.get(n) == t, types.get(n))
        rels1 = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        rels8 = z.read("xl/worksheets/_rels/sheet8.xml.rels").decode()
        check("drawings: the comments' sheet reaches its comments, their boxes and its drawing",
              'Id="rId1" Type="%s/comments"' % RNS in rels1 and 'Id="rId2" Type="%s/vmlDrawing"' % RNS in rels1 and
              '<Relationship Id="rId3" Type="%s/drawing" Target="../drawings/drawing2.xml"/>' % RNS in rels1, rels1)
        check("drawings: the chart's sheet reaches its drawing",
              rels8.count("<Relationship ") == 1 and 'Target="../drawings/drawing1.xml"' in rels8, rels8)
        sheet1 = z.read("xl/worksheets/sheet1.xml").decode()
        sheet8 = z.read("xl/worksheets/sheet8.xml").decode()
        check("drawings: <drawing> after the page margins, before the comments' boxes",
              sheet1.find("<pageMargins") < sheet1.find('<drawing r:id="rId3"/>') < sheet1.find('<legacyDrawing r:id="rId2"/>'),
              sheet1[-300:])
        check("drawings: the chart's sheet's <drawing>", sheet8.endswith('<drawing r:id="rId3"/></worksheet>\n'), sheet8[-200:])
        stored = {i.filename: i.compress_type for i in z.infolist()}
        check("drawings: the picture stored as it is packed, the chart deflated",
              stored.get("xl/media/image1.png") == zipfile.ZIP_STORED and stored.get("xl/charts/chart1.xml") == zipfile.ZIP_DEFLATED,
              stored)
    again = dump(xlsx)
    check("drawings: the same after saving", kept_parts(again) == KEPT and
          lines(again, "drawing") == lines(got, "drawing"), (kept_parts(again), lines(again, "drawing")))
    twice = os.path.join(work, "drawn-twice.xlsx")
    convert(xlsx, twice)
    with zipfile.ZipFile(twice) as z:
        check("drawings: and after saving again", all(z.read(n) == drawn[n] for n in KEPT), "")
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl:
        wb = openpyxl.load_workbook(xlsx)
        charts = wb["Chart"]._charts
        check("drawings: openpyxl reads the chart back", len(charts) == 1 and type(charts[0]).__name__ == "BarChart" and
              (charts[0].anchor._from.col, charts[0].anchor._from.row) == (3, 1), charts)
        try:
            import PIL  # noqa: F401 - openpyxl reads pictures only with it
            check("drawings: openpyxl reads the pictures back", len(wb["Chart"]._images) == 1 and len(wb["Systems"]._images) == 1,
                  (wb["Chart"]._images, wb["Systems"]._images))
        except ImportError:
            pass
    if SOFFICE:
        profile = tempfile.mkdtemp(prefix="owf-lo-")
        try:
            outdir = os.path.join(work, "lo-drawings")
            os.makedirs(outdir, exist_ok=True)
            subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", "ods",
                            "--outdir", outdir, xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=300)
            with zipfile.ZipFile(os.path.join(outdir, "drawn-saved.ods")) as z:
                content = z.read("content.xml").decode()
            check("drawings: LibreOffice reads the chart", "<draw:object " in content and "Chart.B1:Chart.B4" in content,
                  re.findall(r"<draw:object[^>]*>", content))
            check("drawings: LibreOffice reads both pictures", len(re.findall(r'<draw:image xlink:href="Pictures/', content)) == 2,
                  re.findall(r"<draw:image[^>]*>", content))
        except Exception as e:  # noqa: BLE001
            check("drawings: LibreOffice converts the file written", False, e)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
    report = convert(FIX, os.path.join(work, "drawings.ods"))
    check("drawings: ODS says it leaves them out", "Pictures and charts on sheets are kept in Excel workbooks" in report, report)

    # the chart's part missing: that drawing is left out whole (with a note), the other kept
    broken = dict(drawn)
    del broken["xl/charts/chart1.xml"]
    path = os.path.join(work, "drawn-broken.xlsx")
    write_zip(path, broken)
    got = dump(path)
    check("drawings: a drawing with a part missing is left out", sheet_lines(got, 7, "drawing") == [] and
          "could not be kept: their part of the file is missing" in got and
          set(kept_parts(got)) == {"xl/drawings/drawing2.xml", "xl/drawings/_rels/drawing2.xml.rels", "xl/media/image1.png"},
          (sheet_lines(got, 7, "drawing"), kept_parts(got)))
    saved = os.path.join(work, "drawn-broken-saved.xlsx")
    convert(path, saved)
    well_formed(saved)
    with zipfile.ZipFile(saved) as z:
        check("drawings: nothing of it written", "xl/drawings/drawing1.xml" not in z.namelist() and
              "xl/worksheets/_rels/sheet8.xml.rels" not in z.namelist() and "<drawing" not in z.read("xl/worksheets/sheet8.xml").decode(),
              z.namelist())
    # its relationships damaged: left out too
    broken = dict(drawn)
    broken["xl/drawings/_rels/drawing1.xml.rels"] = drawn["xl/drawings/_rels/drawing1.xml.rels"][:60]
    write_zip(path, broken)
    got = dump(path)
    check("drawings: a drawing with damaged relationships is left out", sheet_lines(got, 7, "drawing") == [] and
          "xl/charts/chart1.xml" not in kept_parts(got), kept_parts(got))
    # relationships that go round (a part to itself, the chart back to its drawing), a target outside the package
    looped = dict(drawn)
    looped["xl/charts/_rels/chart1.xml.rels"] = drawn["xl/charts/_rels/chart1.xml.rels"].replace(b"</Relationships>", (
        '<Relationship Id="rId7" Type="%s/image" Target="chart1.xml"/>'
        '<Relationship Id="rId8" Type="%s/image" Target="/xl/drawings/drawing1.xml"/></Relationships>' % (RNS, RNS)).encode())
    write_zip(path, looped)
    got = dump(path)
    check("drawings: relationships that go round keep each part once", set(kept_parts(got)) == set(KEPT), kept_parts(got))
    convert(path, saved)
    with zipfile.ZipFile(saved) as z:
        check("drawings: and write it once", len(z.namelist()) == len(set(z.namelist())), sorted(z.namelist()))
    outside = dict(drawn)
    outside["xl/drawings/_rels/drawing2.xml.rels"] = drawn["xl/drawings/_rels/drawing2.xml.rels"].replace(
        b"/xl/media/image1.png", b"../../../../../../etc/passwd")
    write_zip(path, outside)
    got = dump(path)
    check("drawings: a target outside the package is not followed out of it", sheet_lines(got, 0, "drawing") == [] and
          all(not n.startswith("etc/") for n in kept_parts(got)), kept_parts(got))


# ---- Charts ----

C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
EXCEL_CHART = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<c:chartSpace xmlns:c="%s" xmlns:a="%s" xmlns:r="%s"><c:chart>'
    '<c:title><c:tx><c:strRef><c:f>Chart!$D$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Systems &amp; regions</c:v>'
    '</c:pt></c:strCache></c:strRef></c:tx><c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/><c:plotArea><c:layout/>'
    '<c:barChart><c:barDir val="bar"/><c:grouping val="stacked"/><c:varyColors val="0"/>'
    '<c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:strRef><c:f>Chart!$B$1</c:f><c:strCache><c:ptCount val="1"/>'
    '<c:pt idx="0"><c:v>North</c:v></c:pt></c:strCache></c:strRef></c:tx><c:spPr><a:solidFill><a:schemeClr val="accent2">'
    '<a:lumMod val="75000"/></a:schemeClr></a:solidFill></c:spPr><c:invertIfNegative val="0"/>'
    '<c:cat><c:strRef><c:f>Chart!$A$2:$A$4</c:f><c:strCache><c:ptCount val="3"/><c:pt idx="0"><c:v>Q1</c:v></c:pt>'
    '<c:pt idx="1"><c:v>Q2</c:v></c:pt><c:pt idx="2"><c:v>Q3</c:v></c:pt></c:strCache></c:strRef></c:cat>'
    '<c:val><c:numRef><c:f>Chart!$B$2:$B$4</c:f><c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="3"/>'
    '<c:pt idx="0"><c:v>12</c:v></c:pt><c:pt idx="2"><c:v>-3.5</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser>'
    '<c:ser><c:idx val="3"/><c:order val="1"/><c:tx><c:v>Typed</c:v></c:tx><c:spPr><a:solidFill><a:srgbClr val="00b050"/>'
    '</a:solidFill></c:spPr><c:val><c:numLit><c:ptCount val="2"/><c:pt idx="0"><c:v>1</c:v></c:pt><c:pt idx="1"><c:v>2</c:v>'
    '</c:pt></c:numLit></c:val></c:ser><c:gapWidth val="80"/><c:overlap val="100"/><c:axId val="1"/><c:axId val="2"/></c:barChart>'
    '<c:lineChart><c:grouping val="standard"/><c:ser><c:idx val="9"/><c:val><c:numRef><c:f>Chart!$C$2:$C$4</c:f></c:numRef>'
    '</c:val></c:ser><c:axId val="1"/><c:axId val="2"/></c:lineChart>'
    '<c:catAx><c:axId val="1"/></c:catAx><c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/>'
    '<c:max val="20"/><c:min val="-5"/></c:scaling><c:delete val="0"/></c:valAx></c:plotArea>'
    '<c:legend><c:legendPos val="b"/></c:legend></c:chart></c:chartSpace>' % (C, DML, RNS))


def test_charts(work):
    src = dump(FIX)
    check("charts: the fixture's", sheet_lines(src, 7, "chart") == [
        'chart kind=1 grouping=0 from=1,3+0,0 size=5400000x2700000 legend=r gap=150 grid=1 title "Systems by region"'],
        sheet_lines(src, 7, "chart"))
    check("charts: its series", sheet_lines(src, 7, "series") == [
        "series 0 colour=- name \"\" values \"'Chart'!$B$1:$B$4\" categories \"'Chart'!$A$1:$A$4\" points"],
        sheet_lines(src, 7, "series"))
    with zipfile.ZipFile(FIX) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    # as Excel writes one: a stacked bar chart and a line chart over it (the first group read), a two-cell
    # anchor, its title from a cell, theme and RGB colours, a literal series, gaps, the axis's limits
    p = dict(parts)
    p["xl/charts/chart1.xml"] = EXCEL_CHART.encode()
    p["xl/drawings/drawing1.xml"] = p["xl/drawings/drawing1.xml"].replace(
        b"<oneCellAnchor><from><col>3</col><colOff>0</colOff><row>1</row><rowOff>0</rowOff></from>"
        b'<ext cx="5400000" cy="2700000"/>',
        b"<twoCellAnchor><from><col>3</col><colOff>95250</colOff><row>1</row><rowOff>9525</rowOff></from>"
        b"<to><col>10</col><colOff>0</colOff><row>16</row><rowOff>190500</rowOff></to>").replace(
        b"</oneCellAnchor>", b"</twoCellAnchor>")
    path = os.path.join(work, "excel-chart.xlsx")
    write_zip(path, p)
    got = dump(path)
    check("charts: Excel's", sheet_lines(got, 7, "chart") == [
        'chart kind=2 grouping=1 from=1,3+9525,95250 to=16,10+190500,0 legend=b gap=80 grid=0 min=-5 max=20 '
        'title "Systems & regions"'], sheet_lines(got, 7, "chart"))
    check("charts: Excel's series", sheet_lines(got, 7, "series") == [
        'series 0 colour=ED7D31 name "North" from "Chart!$B$1" values "Chart!$B$2:$B$4" categories "Chart!$A$2:$A$4" '
        'points 12="Q1" -="Q2" -3.5="Q3"',
        'series 3 colour=00B050 name "Typed" values "" categories "" points 1 2'], sheet_lines(got, 7, "series"))
    # a picture alone, a chart part missing or damaged: no chart, no error
    for name, data in (("xl/charts/chart1.xml", EXCEL_CHART[:400].encode()), ("xl/charts/chart1.xml", b"")):
        bad = dict(p, **{name: data})
        write_zip(path, bad)
        rc, out, err = run([DUMP, path])
        check("charts: a damaged chart read safely", rc == 0 and not sheet_lines(out, 7, "chart"), out[-300:] + err[-300:])
    # the charts fixture: lines, stacked bars, 100% columns (legend below), a pie on another sheet; the same after saving
    got = dump(CHARTS)
    want = [('chart kind=3 grouping=0 from=1,5+0,0 size=5400000x2700000 legend=r gap=150 grid=1 title "Sales by month"', 0),
            ('chart kind=2 grouping=1 from=17,5+0,0 size=5400000x2700000 legend=r gap=150 grid=1 title "Stacked bars"', 0),
            ('chart kind=1 grouping=2 from=1,15+0,0 size=5400000x2700000 legend=b gap=150 grid=1 title "Share"', 0),
            ('chart kind=5 grouping=0 from=1,1+0,0 size=5400000x2700000 legend=r gap=150 grid=0 title "Amiga by month"', 1)]
    check("charts: the charts fixture's", sheet_lines(got, 0, "chart") == [w for w, sh in want if sh == 0] and
          sheet_lines(got, 1, "chart") == [w for w, sh in want if sh == 1], lines(got, "chart"))
    check("charts: its series' names from their cells", lines(got, "series")[:3] == [
        'series %d colour=- name "" from "\'Sales\'!%s1" values "\'Sales\'!$%s$2:$%s$7" categories "\'Sales\'!$A$2:$A$7" points'
        % (k, c, c, c) for k, c in enumerate("BCD")], lines(got, "series")[:3])
    saved = os.path.join(work, "charts-saved.xlsx")
    convert(CHARTS, saved)
    again = dump(saved)
    check("charts: the same after saving", lines(again, "chart") == lines(got, "chart") and
          lines(again, "series") == lines(got, "series"), lines(again, "chart"))
    pics = drawn_book(parts)
    write_zip(path, pics)
    got = dump(path)
    check("charts: pictures are not charts", sheet_lines(got, 0, "chart") == [] and len(sheet_lines(got, 7, "chart")) == 1,
          (sheet_lines(got, 0, "chart"), sheet_lines(got, 7, "chart")))


def damaged(work):
    """The first sheet's XML cut short and changed in what follows its cells (the conditional formats and
    the rest): book_dump answers with a book or an error, never a crash."""
    with zipfile.ZipFile(FIX) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"]
    start = sheet.find(b"</sheetData>")
    rnd = __import__("random").Random(25102026)
    bad = 0
    for k in range(60):
        cut = start + (len(sheet) - start) * k // 60
        variants = [sheet[:cut]]
        if cut < len(sheet):
            changed = bytearray(sheet)
            changed[cut] = rnd.choice(b'<>"=/ x')
            variants.append(bytes(changed))
        for v in variants:
            path = os.path.join(work, "damaged.xlsx")
            with zipfile.ZipFile(path, "w") as z:
                for n, data in parts.items():
                    z.writestr(n, v if n == "xl/worksheets/sheet1.xml" else data)
            rc, out, err = run([DUMP, path])
            if rc not in (0, 10):
                bad += 1
                check("damaged sheet read safely (cut at %d)" % cut, False, err[-800:])
    check("damaged sheets read safely", bad == 0, bad)
    # the drawings' relationships and content types cut short or changed: read, and written, safely
    drawn = drawn_book(parts)
    bad = 0
    drawn["xl/charts/chart1.xml"] = EXCEL_CHART.encode()
    for name in ("xl/drawings/_rels/drawing1.xml.rels", "xl/charts/_rels/chart1.xml.rels", "[Content_Types].xml",
                 "xl/worksheets/_rels/sheet8.xml.rels", "xl/charts/chart1.xml", "xl/drawings/drawing1.xml"):
        data = drawn[name]
        for k in range(0, len(data), max(1, len(data) // 25)):
            changed = bytearray(data)
            changed[k] = rnd.choice(b'<>"=/ x')
            for v in (data[:k], bytes(changed)):
                path = os.path.join(work, "damaged-drawing.xlsx")
                write_zip(path, dict(drawn, **{name: v}))
                rc, out, err = run([DUMP, path])
                if rc not in (0, 10):
                    bad += 1
                    check("damaged drawing read safely (%s at %d)" % (name, k), False, err[-800:])
                elif rc == 0:
                    rc, out, err = run([OWCONVERT, path, os.path.join(work, "damaged-drawing-saved.xlsx")])
                    if rc not in (0, 10):
                        bad += 1
                        check("damaged drawing written safely (%s at %d)" % (name, k), False, err[-800:])
    check("damaged drawings read and written safely", bad == 0, bad)


def main():
    with tempfile.TemporaryDirectory() as work:
        test_condfmt(work)
        test_comments(work)
        test_validation(work)
        test_autofilter(work)
        test_drawings(work)
        test_charts(work)
        damaged(work)
    if failures:
        print("sheet features: %d of %d checks failed" % (len(failures), checks))
        return 1
    print("sheet features: %d checks passed" % checks)
    return 0


if __name__ == "__main__":
    sys.exit(main())
