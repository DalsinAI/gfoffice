#!/usr/bin/env python3
"""Makes the XLSX and ODS test files: make_xml_fixtures.py [OUTDIR]

Files other programs wrote, for test_xml_books.py:
  book-lo.xlsx, book-lo.ods   LibreOffice's saves of book.fods (the sheets
                              make_fixtures.py writes), formulas worked out
  book-openpyxl.xlsx          openpyxl's, with what it can write that
                              LibreOffice's save of book.fods has not
  book-odfpy.ods              odfpy's, the same for ODS
openpyxl and odfpy (pip install openpyxl odfpy) and LibreOffice only make
these files; nothing of them is in libowf. Run it again after changing the
sheets below, and change test_xml_books.py to match.

The names of Excel's functions in OpenFormula that owf_xlformula.c uses
were found the same way: a workbook calling each function, saved by
LibreOffice as ODS (functions_odf() below writes it; see --functions).
MIT, Copyright (c) 2026 AmigaChrome Team.
"""
import datetime
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else os.path.join(HERE, "fixtures"))

RECALC = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def soffice(source, filter_name, outdir):
    """LibreOffice converts source, with a profile of its own that works
    the formulas out."""
    profile = tempfile.mkdtemp(prefix="owf-lo-")
    try:
        os.makedirs(os.path.join(profile, "user"))
        with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as f:
            f.write(RECALC)
        subprocess.run(["soffice", "-env:UserInstallation=file://" + profile, "--headless",
                        "--convert-to", filter_name, "--outdir", outdir, source],
                       check=True, stdout=subprocess.DEVNULL, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def libreoffice_books():
    source = os.path.join(HERE, "fixtures", "book.fods")
    with tempfile.TemporaryDirectory() as work:
        soffice(source, 'xlsx:Calc MS Excel 2007 XML', work)
        soffice(source, "ods", work)
        shutil.copy(os.path.join(work, "book.xlsx"), os.path.join(OUT, "book-lo.xlsx"))
        shutil.copy(os.path.join(work, "book.ods"), os.path.join(OUT, "book-lo.ods"))


def openpyxl_book():
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.workbook.defined_name import DefinedName
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
    except ImportError:
        CellRichText = None

    wb = openpyxl.Workbook()
    wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Sums"
    ws["A1"] = "Name"
    ws["A1"].font = Font(name="Verdana", size=14, bold=True, italic=True, color="FF336699")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFFCC00")
    ws["A1"].border = Border(left=Side("thin", color="FF00FF00"), bottom=Side("double"))
    ws["A1"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True, indent=1)
    for row, (name, qty) in enumerate([("one", 1), ("two", 2.5), ("three", -3)], start=2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=qty)
        ws.cell(row=row, column=3, value="=B%d*2" % row)
    ws["B5"] = "=SUM(B2:B4)"
    ws["C5"] = "=_xlfn.CONCAT(A2,A3)"
    ws["D2"] = ArrayFormula("D2:D4", "=B2:B4*10")
    ws["E2"] = True
    ws["E3"] = "#N/A"
    ws["E3"].data_type = "e"
    ws["F2"] = datetime.datetime(2026, 10, 25, 1, 30)
    ws["F2"].number_format = "yyyy-mm-dd hh:mm"
    ws["F3"] = 0.125
    ws["F3"].number_format = "0.0%"
    ws["F4"] = 1234567.891
    ws["F4"].number_format = '#,##0.00 "EUR";[Red]-#,##0.00 "EUR"'
    ws["G2"] = "Text\twith tab and _x0041_"
    if CellRichText:
        ws["G3"] = CellRichText("plain ", TextBlock(InlineFont(b=True), "bold"))
    ws.merge_cells("A7:C8")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 18.5
    ws.column_dimensions["H"].hidden = True
    ws.row_dimensions[6].height = 33
    ws.row_dimensions[9].hidden = True
    ws["A1"].comment = Comment("a note", "me")
    ws.conditional_formatting.add("B2:B4", CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill("solid", fgColor="FFFF0000")))
    dv = DataValidation(type="list", formula1='"a,b,c"')
    ws.add_data_validation(dv)
    dv.add("A10")
    other = wb.create_sheet("Other sheet")
    other["A1"] = "='Sums'!B5+1"
    other["A2"] = "=Total*2"
    other.sheet_state = "hidden"
    wb.defined_names["Total"] = DefinedName("Total", attr_text="Sums!$B$5")
    third = wb.create_sheet("Third")
    third["B2"] = 7
    wb.active = 2
    wb.save(os.path.join(OUT, "book-openpyxl.xlsx"))


def odfpy_book():
    from odf.namespaces import TABLENS
    from odf.number import Day, Month, NumberStyle, Number, DateStyle, Text as NText, Year, PercentageStyle
    from odf.office import Annotation
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.style import (FontFace, ParagraphProperties, Style, TableCellProperties, TableColumnProperties,
                           TableRowProperties, TextProperties)
    from odf.table import (CoveredTableCell, NamedExpressions, NamedRange, Table, TableCell, TableColumn,
                           TableRow, CalculationSettings, NullDate)
    from odf.text import P, S, Span, LineBreak, Tab

    doc = OpenDocumentSpreadsheet()
    doc.fontfacedecls.addElement(FontFace(name="DejaVu Serif", fontfamily="'DejaVu Serif'", fontfamilygeneric="roman"))
    money = NumberStyle(name="N1")
    money.addElement(Number(decimalplaces="2", minintegerdigits="1", grouping="true"))
    money.addElement(NText(text=" kg"))
    doc.automaticstyles.addElement(money)
    pct = PercentageStyle(name="N2")
    pct.addElement(Number(decimalplaces="1", minintegerdigits="1"))
    pct.addElement(NText(text="%"))
    doc.automaticstyles.addElement(pct)
    date = DateStyle(name="N3")
    date.addElement(Day(style="long"))
    date.addElement(NText(text="."))
    date.addElement(Month(style="long"))
    date.addElement(NText(text="."))
    date.addElement(Year(style="long"))
    doc.automaticstyles.addElement(date)
    head = Style(name="ce1", family="table-cell", datastylename="N1")
    head.addElement(TableCellProperties(backgroundcolor="#ccffcc", border="0.06pt solid #ff0000", wrapoption="wrap",
                                        verticalalign="middle"))
    head.addElement(ParagraphProperties(textalign="end", marginleft="0.2in"))
    head.addElement(TextProperties(fontname="DejaVu Serif", fontsize="13pt", fontweight="bold", color="#0000ff",
                                   textlinethroughstyle="solid"))
    doc.automaticstyles.addElement(head)
    for name, ds in (("ce2", "N2"), ("ce3", "N3")):
        doc.automaticstyles.addElement(Style(name=name, family="table-cell", datastylename=ds))
    wide = Style(name="co1", family="table-column")
    wide.addElement(TableColumnProperties(columnwidth="1.5in"))
    doc.automaticstyles.addElement(wide)
    tall = Style(name="ro1", family="table-row")
    tall.addElement(TableRowProperties(rowheight="0.5in", useoptimalrowheight="false"))
    doc.automaticstyles.addElement(tall)

    settings = CalculationSettings()
    settings.addElement(NullDate(datevalue="1904-01-01"))
    doc.spreadsheet.addElement(settings)
    t = Table(name="Odf one")
    t.addElement(TableColumn(stylename="co1"))
    t.addElement(TableColumn(numbercolumnsrepeated="2"))
    t.addElement(TableColumn(visibility="collapse"))
    rows = [
        [("Head", "string", "ce1"), (1.5, "float", "ce1"), (0.25, "percentage", "ce2"), ("2026-10-21", "date", "ce3")],
        [(2, "float", None), (3, "float", None), ("of:=SUM([.A2:.B2])", 5, None), ("of:=IF([.A2]>1;\"big\";\"small\")", "big", None)],
        [(True, "boolean", None), ("PT36H00M00S", "time", None), ("of:=[$'Odf two'.A1]*2", 84, None), ("of:=COM.MICROSOFT.CEILING([.A2];3)", 3, None)],
    ]
    for i, cells in enumerate(rows):
        tr = TableRow(stylename="ro1") if i == 0 else TableRow()
        for value, kind, style in cells:
            kw = {"stylename": style} if style else {}
            if isinstance(value, str) and value.startswith("of:="):
                if isinstance(kind, str):
                    tc = TableCell(formula=value, valuetype="string", stringvalue=kind, **kw)
                    tc.addElement(P(text=kind))
                else:
                    tc = TableCell(formula=value, valuetype="float", value=kind, **kw)
                    tc.addElement(P(text=str(kind)))
            elif kind == "string":
                tc = TableCell(valuetype="string", **kw)
                tc.addElement(P(text=value))
            elif kind == "boolean":
                tc = TableCell(valuetype="boolean", booleanvalue="true" if value else "false", **kw)
            elif kind == "date":
                tc = TableCell(valuetype="date", datevalue=value, **kw)
            elif kind == "time":
                tc = TableCell(valuetype="time", timevalue=value, **kw)
            else:
                tc = TableCell(valuetype=kind, value=value, **kw)
            tr.addElement(tc)
        t.addElement(tr)
    # a cell of text in pieces: spaces, a tab, a line break, two paragraphs, a span
    tr = TableRow()
    tc = TableCell(valuetype="string", numbercolumnsspanned="2", numberrowsspanned="1")
    p = P(text="a")
    p.addElement(S(c=3))
    p.addText("b")
    p.addElement(Tab())
    p.addText("c")
    p.addElement(LineBreak())
    span = Span(text="d")
    p.addElement(span)
    tc.addElement(p)
    tc.addElement(P(text="second"))
    ann = Annotation()
    ann.addElement(P(text="a comment"))
    tc.addElement(ann)
    tr.addElement(tc)
    tr.addElement(CoveredTableCell())
    tr.addElement(TableCell(numbercolumnsrepeated="2"))
    t.addElement(tr)
    t.addElement(TableRow(numberrowsrepeated="3", visibility="collapse"))
    tr = TableRow()
    tr.addElement(TableCell(valuetype="float", value="7", numbercolumnsrepeated="3"))
    t.addElement(tr)
    doc.spreadsheet.addElement(t)
    t2 = Table(name="Odf two")
    tr = TableRow()
    tc = TableCell(valuetype="float", value="42")
    tc.addElement(P(text="42"))
    tr.addElement(tc)
    t2.addElement(tr)
    doc.spreadsheet.addElement(t2)
    ne = NamedExpressions()
    ne.addElement(NamedRange(name="Two", cellrangeaddress="$'Odf two'.$A$1", basecelladdress="$'Odf two'.$A$1"))
    doc.spreadsheet.addElement(ne)
    doc.save(os.path.join(OUT, "book-odfpy.ods"))


def functions_odf(names, path):
    """--functions NAME...: an XLSX calling each function, for LibreOffice
    to save as ODS; how each name comes out is what owf_xlformula.c's
    tables say."""
    import openpyxl
    wb = openpyxl.Workbook()
    for i, n in enumerate(names, start=1):
        wb.active.cell(row=i, column=1, value="=%s(1)" % n)
    wb.save(path)


def main():
    if "--functions" in sys.argv:
        names = sys.argv[sys.argv.index("--functions") + 1:]
        with tempfile.TemporaryDirectory() as work:
            src = os.path.join(work, "f.xlsx")
            functions_odf(names, src)
            soffice(src, "fods", work)
            print(open(os.path.join(work, "f.fods"), encoding="utf-8").read())
        return
    os.makedirs(OUT, exist_ok=True)
    libreoffice_books()
    openpyxl_book()
    odfpy_book()
    for name in ("book-lo.xlsx", "book-lo.ods", "book-openpyxl.xlsx", "book-odfpy.ods"):
        print(os.path.join(OUT, name), os.path.getsize(os.path.join(OUT, name)))


if __name__ == "__main__":
    main()
