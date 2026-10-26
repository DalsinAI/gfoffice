#!/usr/bin/env python3
"""Makes fixtures/features-openpyxl.xlsx, fixtures/charts-openpyxl.xlsx and fixtures/validation-openpyxl.xlsx:
make_feature_fixtures.py [OUTDIR] [validation]

A workbook made up to have what the lab's yardstick has (a team's
systems list, 25 October 2026, which stays private and is never put
here): eight sheets, one of them empty; on the first, about 240 rows of
text, numbers and formulas under a dark header row, frozen panes at B7,
merged ranges, conditional formatting (cellIs, expression, colour scales,
a data bar and a "contains text" rule), data validation, an AutoFilter,
comments, grouped rows, column widths and row heights as Excel writes
them (not custom), cells aligned top, centre and bottom, theme colours
and custom number formats; a bar chart on another sheet. And a workbook
of charts OpenCalc draws: lines, stacked bars and 100% columns over a
table with frozen panes; a pie (one slice coloured by hand) and a
doughnut; a scatter of two series, one markers only; data labels (values
on columns, values above a line's markers, categories and shares on a
pie) and axes' number formats (one set on the axis, one the cells'). And
a workbook of validation rules that alert, one to a row, to type into in
a lab: lists (written out, from cells, from a name, without their arrow,
without an alert), numbers, a date, a time, a text's length and a
formula, with Stop, Warning and Information alerts. "validation" makes
only that one.
test_sheet_features.py checks what libowf reads of it and keeps on
saving. openpyxl (pip install openpyxl) only makes the file; nothing of it
is in libowf. Run it again after changing the sheets below, and change
test_sheet_features.py to match.
MIT, Copyright (c) 2026 AmigaChrome Team.
"""
import datetime
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "fixtures"))

SYSTEMS = ["Directory", "Mail relay", "Name service", "Content network", "Firewall", "Edge cache", "Build service",
           "Code host", "Code assistant", "Code checks", "Ticket desk", "Wiki", "Chat", "File share", "Backups"]
KINDS = ["IaaS/PaaS - public cloud", "PaaS - customer identity", "PaaS - DNS hosting", "PaaS - content network",
         "Self-hosted software - static analysis", "SaaS - source control"]


def systems_sheet(ws):
    thin = Side(style="thin", color=Color(theme=1, tint=0.5))
    medium = Side(style="medium", color="FF000000")
    ws["A2"] = "Systems - Critical Systems List"
    ws["A2"].font = Font(name="Arial", size=14, bold=True, underline="single")
    ws["A3"] = "High level critical systems and suppliers"
    ws["A3"].font = Font(name="Arial", size=11, bold=True, color=Color(theme=4))
    ws.merge_cells("A2:C2")
    ws.merge_cells("A3:C3")
    ws.merge_cells("D5:F5")
    ws["D5"] = "Figures"
    ws["D5"].alignment = Alignment(horizontal="right")
    heads = ["Application", "Version", "Software/Service Type", "Hosted", "Owner", "Cost", "Renewal", "Score", "Notes"]
    dark = PatternFill("solid", fgColor=Color(theme=1, tint=0.25))
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.font = Font(name="Arial", size=11, bold=True, color=Color(theme=0))
        cell.fill = dark
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    for r in range(7, 243):
        i = r - 7
        ws.cell(row=r, column=1, value=SYSTEMS[i % len(SYSTEMS)] + ("" if i < len(SYSTEMS) else " %d" % (i // len(SYSTEMS))))
        ws.cell(row=r, column=1).alignment = Alignment(vertical="bottom")
        ws.cell(row=r, column=2, value="N/A" if i % 4 else "To confirm").alignment = Alignment(vertical="top")
        ws.cell(row=r, column=3, value=KINDS[i % len(KINDS)]).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=4, value="Yes" if i % 3 else "No").alignment = Alignment(vertical="center", horizontal="center")
        ws.cell(row=r, column=5, value="Team %d" % (i % 5 + 1))
        ws.cell(row=r, column=6, value=round(100 + (i * 37) % 900 + 0.25 * (i % 4), 2)).number_format = '"£"#,##0.00'
        ws.cell(row=r, column=7, value=datetime.date(2026, 1, 1) + datetime.timedelta(days=(i * 11) % 365)).number_format = "dd mmm yyyy"
        ws.cell(row=r, column=8, value="=IF(D%d=\"Yes\",F%d/10,0)" % (r, r)).number_format = "0.0"
        if i % 7 == 0:
            ws.cell(row=r, column=9, value="Review due")
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.row_dimensions[r].height = 14.25
        if 20 <= r <= 30:
            ws.row_dimensions[r].outlineLevel = 1
    ws.row_dimensions[2].height = 25.5
    ws.row_dimensions[3].height = 15
    ws.row_dimensions[5].height = 21.75
    ws.row_dimensions[6].height = 27
    for col, width in zip("ABCDEFGHI", (45.28, 22, 31.42, 17.42, 35.57, 29.71, 14, 10, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B7"
    ws.auto_filter.ref = "A6:I242"
    # conditional formatting: cellIs, expression, colour scales, a data bar, contains text
    red = DifferentialStyle(font=Font(bold=True, color="FF9C0006"), fill=PatternFill(bgColor="FFFFC7CE"))
    ws.conditional_formatting.add("F7:F242", CellIsRule(operator="greaterThan", formula=["800"], fill=PatternFill(bgColor="FFFFC7CE"),
                                                       font=Font(color="FF9C0006")))
    ws.conditional_formatting.add("F7:F242", CellIsRule(operator="between", formula=["100", "200"], fill=PatternFill(bgColor="FFC6EFCE")))
    ws.conditional_formatting.add("A7:I242", FormulaRule(formula=['$D7="No"'], font=Font(italic=True, color="FF808080")))
    ws.conditional_formatting.add("H7:H242", ColorScaleRule(start_type="min", start_color="FFF8696B", mid_type="percentile",
                                                           mid_value=50, mid_color="FFFFEB84", end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("G7:G242", ColorScaleRule(start_type="min", start_color="FFFFFFFF", end_type="max",
                                                           end_color="FF5A8AC6"))
    ws.conditional_formatting.add("F7:F242", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    text_rule = Rule(type="containsText", operator="containsText", text="Review", dxf=red)
    text_rule.formula = ['NOT(ISERROR(SEARCH("Review",I7)))']
    ws.conditional_formatting.add("I7:I242", text_rule)
    stop_rule = CellIsRule(operator="equal", formula=['"To confirm"'], fill=PatternFill(bgColor="FFFFEB9C"), stopIfTrue=True)
    ws.conditional_formatting.add("B7:B242", stop_rule)
    ws.conditional_formatting.add("E7:E242", FormulaRule(formula=["LEN(E7)>5"], fill=PatternFill(bgColor="FFDDEBF7")))
    ws.conditional_formatting.add("C7:C242", CellIsRule(operator="notEqual", formula=['""'], font=Font(bold=True)))
    # data validation
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv.prompt, dv.promptTitle, dv.error, dv.errorTitle = "Hosted by us?", "Hosted", "Yes or No", "Not a choice"
    ws.add_data_validation(dv)
    dv.add("D7:D242")
    dv2 = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000")
    ws.add_data_validation(dv2)
    dv2.add("F7:F242")
    dv3 = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)")
    ws.add_data_validation(dv3)
    dv3.add("G7:G242")
    dv4 = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="200", showErrorMessage=True, errorStyle="warning")
    ws.add_data_validation(dv4)
    dv4.add("I7:I242")
    dv5 = DataValidation(type="list", formula1="Lists!$A$1:$A$5")
    ws.add_data_validation(dv5)
    dv5.add("E7:E242")
    dv6 = DataValidation(type="whole", operator="between", formula1="0", formula2="10")
    ws.add_data_validation(dv6)
    dv6.add("H7:H20")
    dv6.add("H30:H40")
    dv7 = DataValidation(type="custom", formula1="LEN(A7)>0")
    ws.add_data_validation(dv7)
    dv7.add("A7:A242")
    # comments (written as legacy VML drawings)
    ws["A7"].comment = Comment("Main identity service.\nOwner to confirm.", "Team 5")
    ws["C9"].comment = Comment("Hosted by the supplier", "Team 5")
    ws["F12"].comment = Comment("Cost per year, before VAT", "Finance")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Systems"
    systems_sheet(ws)
    rep = wb.create_sheet("Reporting")
    for r in range(1, 120):
        rep.cell(row=r, column=1, value="Item %d" % r)
        rep.cell(row=r, column=2, value=r * 3)
        rep.cell(row=r, column=3, value="=B%d*2" % r)
        rep.cell(row=r, column=4, value="=SUM($B$1:B%d)" % r)
    cost = wb.create_sheet("Cost")
    for r in range(1, 86):
        cost.cell(row=r, column=1, value="Line %d" % r)
        cost.cell(row=r, column=2, value=r * 12.5)
        cost.cell(row=r, column=3, value="=B%d*1.2" % r)
        cost.cell(row=r, column=4, value="=C%d-B%d" % (r, r))
    log = wb.create_sheet("Log")
    for r in range(1, 12):
        log.cell(row=r, column=1, value="Entry %d" % r)
        log.cell(row=r, column=2, value="=LEN(A%d)" % r)
        log.cell(row=r, column=3, value="=B%d+%d" % (r, r))
    wb.create_sheet("Empty")
    lists = wb.create_sheet("Lists")
    for r, v in enumerate(["Team 1", "Team 2", "Team 3", "Team 4", "Team 5"], 1):
        lists.cell(row=r, column=1, value=v)
    gaps = wb.create_sheet("Gaps")
    for r in range(1, 30):
        gaps.cell(row=r, column=1, value="Gap %d" % r)
        gaps.cell(row=r, column=2, value=(r * 7) % 11)
    chart_sheet = wb.create_sheet("Chart")
    for r, (name, value) in enumerate([("North", 12), ("South", 7), ("East", 9), ("West", 15)], 1):
        chart_sheet.cell(row=r, column=1, value=name)
        chart_sheet.cell(row=r, column=2, value=value)
    chart = BarChart()
    chart.title = "Systems by region"
    chart.add_data(Reference(chart_sheet, min_col=2, min_row=1, max_row=4), titles_from_data=False)
    chart.set_categories(Reference(chart_sheet, min_col=1, min_row=1, max_row=4))
    chart_sheet.add_chart(chart, "D2")
    os.makedirs(OUT, exist_ok=True)
    wb.save(os.path.join(OUT, "features-openpyxl.xlsx"))
    make_charts()


def make_charts():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Amiga", "Atari", "PC"])
    for row in [("Jan", 12, 8, 30), ("Feb", 15, 9, 28), ("Mar", 9, 12, 35), ("Apr", 18, 7, 40), ("May", 22, 10, 38),
                ("Jun", 25, 6, 45)]:
        ws.append(row)
    ws.freeze_panes = "B2"

    def series(chart):
        chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=7), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
        return chart
    lines = series(LineChart())
    lines.title = "Sales by month"
    ws.add_chart(lines, "F2")
    bars = series(BarChart())
    bars.type = "bar"
    bars.grouping = "stacked"
    bars.overlap = 100
    bars.title = "Stacked bars"
    ws.add_chart(bars, "F18")
    share = series(BarChart())
    share.grouping = "percentStacked"
    share.overlap = 100
    share.title = "Share"
    share.legend.position = "b"
    ws.add_chart(share, "P2")
    pie_sheet = wb.create_sheet("Pie")
    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
    pie.title = "Amiga by month"
    slice3 = DataPoint(idx=2)
    slice3.graphicalProperties.solidFill = "C00000"
    pie.series[0].dPt.append(slice3)
    pie_sheet.add_chart(pie, "B2")
    ring = DoughnutChart()
    ring.add_data(Reference(ws, min_col=3, max_col=4, min_row=1, max_row=7), titles_from_data=True)
    ring.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
    ring.title = "Atari and PC"
    ring.holeSize = 60
    pie_sheet.add_chart(ring, "L2")
    xy = wb.create_sheet("Scatter")
    xy.append(["MHz", "Speed", "Price (100s)"])
    for row in [(7, 1.0, 5), (14, 2.1, 9), (25, 4.4, 15), (33, 6.0, 18), (50, 9.5, 26)]:
        xy.append(row)
    scatter = ScatterChart()
    scatter.title = "Speed and price by MHz"
    scatter.style = 13
    xs = Reference(xy, min_col=1, min_row=2, max_row=6)
    speed = Series(Reference(xy, min_col=2, min_row=1, max_row=6), xs, title_from_data=True)
    speed.marker.symbol = "circle"
    price = Series(Reference(xy, min_col=3, min_row=1, max_row=6), xs, title_from_data=True)
    price.marker.symbol = "diamond"
    price.graphicalProperties.line.noFill = True
    scatter.series.append(speed)
    scatter.series.append(price)
    xy.add_chart(scatter, "E2")
    lab = wb.create_sheet("Labels")
    lab.append(["Quarter", "Cost", "Visits"])
    for row in [("Q1", 1250, 40), ("Q2", 2400, 55), ("Q3", 1800, 47), ("Q4", 3100, 62)]:
        lab.append(row)
    for r in range(2, 6):
        lab.cell(row=r, column=2).number_format = '"£"#,##0'
    cost = BarChart()
    cost.title = "Cost by quarter"
    cost.add_data(Reference(lab, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    cost.set_categories(Reference(lab, min_col=1, min_row=2, max_row=5))
    cost.dataLabels = DataLabelList()
    cost.dataLabels.showVal = True
    cost.dataLabels.position = "outEnd"
    cost.legend = None
    lab.add_chart(cost, "E2")
    visits = LineChart()
    visits.title = "Visits"
    visits.add_data(Reference(lab, min_col=3, min_row=1, max_row=5), titles_from_data=True)
    visits.set_categories(Reference(lab, min_col=1, min_row=2, max_row=5))
    visits.series[0].marker.symbol = "square"
    visits.series[0].marker.size = 7
    visits.series[0].dLbls = DataLabelList()
    visits.series[0].dLbls.showVal = True
    visits.series[0].dLbls.position = "t"
    visits.y_axis.number_format = "0.0"
    visits.y_axis.numFmt.sourceLinked = False
    visits.legend = None
    lab.add_chart(visits, "E18")
    share = PieChart()
    share.title = "Share of cost"
    share.add_data(Reference(lab, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    share.set_categories(Reference(lab, min_col=1, min_row=2, max_row=5))
    share.dataLabels = DataLabelList()
    share.dataLabels.showCatName = True
    share.dataLabels.showPercent = True
    lab.add_chart(share, "O2")
    wb.save(os.path.join(OUT, "charts-openpyxl.xlsx"))


# What to type into column B, row by row: (what A says, the rule's arguments)
CHECKS = [
    ("Yes or No (list, Stop)",
     dict(type="list", formula1='"Yes,No"', errorStyle="stop", errorTitle="Not a choice", error="Yes or No, please.")),
    ("A team, from Teams!A1:A5 (list, Stop, the default alert)",
     dict(type="list", formula1="Teams!$A$1:$A$5", errorStyle="stop")),
    ("A team, by the name TeamNames (list, Stop)",
     dict(type="list", formula1="TeamNames", errorStyle="stop", error="One of the teams.")),
    ("A score, 1 to 10 (whole number, Warning)",
     dict(type="whole", operator="between", formula1="1", formula2="10", errorStyle="warning", errorTitle="Score",
          error="Scores run from 1 to 10.")),
    ("A share, 0 to 1 (decimal, Stop)",
     dict(type="decimal", operator="between", formula1="0", formula2="1", errorStyle="stop")),
    ("A date from 1 Jan 2020 (date, Information)",
     dict(type="date", operator="greaterThanOrEqual", formula1="DATE(2020,1,1)", errorStyle="information",
          errorTitle="Date", error="Dates before 2020 are unusual here.")),
    ("A time, 09:00 to 17:00 (time, Stop)",
     dict(type="time", operator="between", formula1="TIME(9,0,0)", formula2="TIME(17,0,0)", errorStyle="stop")),
    ("Up to 10 characters (text length, Stop)",
     dict(type="textLength", operator="lessThanOrEqual", formula1="10", errorStyle="stop", error="10 characters at most.")),
    ("No more than the score in C5 (formula, Stop)",
     dict(type="custom", formula1="B10<=$C$5", errorStyle="stop", error="Not more than C5.")),
    ("Yes or No, without its arrow (list, Stop)",
     dict(type="list", formula1='"Yes,No"', errorStyle="stop", showDropDown=True)),
    ("Red, Green or Blue, anything allowed (list, no alert)",
     dict(type="list", formula1='"Red,Green,Blue"', showErrorMessage=False)),
]


def make_validation():
    """validation-openpyxl.xlsx: CHECKS, a rule to a row in B2:B12, B and C alike (C2:C12 already filled in
    with a value each rule takes), and the teams on their own sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Checks"
    ws["A1"], ws["B1"], ws["C1"] = "What to type into B", "Type here", "Takes"
    for cell in (ws["A1"], ws["B1"], ws["C1"]):
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    takes = ["Yes", "Team 2", "Team 5", 7, 0.25, datetime.date(2024, 3, 1), datetime.time(10, 30), "Short", 4,
             "No", "Purple"]
    for i, ((what, rule), value) in enumerate(zip(CHECKS, takes)):
        row = i + 2
        ws.cell(row=row, column=1, value=what)
        ws.cell(row=row, column=3, value=value)
        if isinstance(value, datetime.date):
            ws.cell(row=row, column=3).number_format = "d mmm yyyy"
        elif isinstance(value, datetime.time):
            ws.cell(row=row, column=3).number_format = "hh:mm"
        args = dict(rule)
        args.setdefault("showErrorMessage", True)
        args.setdefault("allow_blank", True)
        dv = DataValidation(**args)
        ws.add_data_validation(dv)
        dv.add("B%d:C%d" % (row, row))
    teams = wb.create_sheet("Teams")
    for r, v in enumerate(["Team 1", "Team 2", "Team 3", "Team 4", "Team 5"], 1):
        teams.cell(row=r, column=1, value=v)
    wb.defined_names["TeamNames"] = DefinedName("TeamNames", attr_text="Teams!$A$1:$A$5")
    os.makedirs(OUT, exist_ok=True)
    wb.save(os.path.join(OUT, "validation-openpyxl.xlsx"))


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[2] == "validation":
        make_validation()
    else:
        main()
        make_validation()
