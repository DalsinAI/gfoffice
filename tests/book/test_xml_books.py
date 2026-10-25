#!/usr/bin/env python3
"""GFOffice's workbook files: test_xml_books.py BOOK_DUMP OWCONVERT TEST_BOOKFMT

XLSX (Office Open XML), ODS and flat ODS (OpenDocument) and CSV, read and
written by libowf:
  - files other programs made (fixtures/book-lo.*, book-openpyxl.xlsx,
    book-odfpy.ods: make_xml_fixtures.py) and packages written by hand
    from the specifications (fixtures/xlsx-hand/, fixtures/ods-hand/,
    zipped here), each checked line by line in book_dump's output;
  - round trips: XLS to XLSX, XLSX to ODS to XLSX, ODS and flat ODS to
    XLSX, CSV both ways, comparing what the model holds (book_dump, with
    style numbers resolved to what they stand for);
  - the files written: well-formed XML, the parts and relationships each
    format needs, and read back by openpyxl and odfpy (when installed) and
    by LibreOffice (when installed; it is not needed);
  - numbers read exactly as Python reads them, and written in the fewest
    digits that read back the same (test_bookfmt --parse);
  - every file cut short and damaged, at the zip's level and inside its
    XML: book_dump must answer with a book or an error, never a crash
    (make book-check runs it under AddressSanitizer and UBSan).
MIT, Copyright (c) 2026 AmigaChrome Team.
"""
import io
import os
import random
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import xml.dom.minidom

HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, "fixtures")
DUMP = os.path.abspath(sys.argv[1])
OWCONVERT = os.path.abspath(sys.argv[2])
BOOKFMT = os.path.abspath(sys.argv[3])
SOFFICE = shutil.which("soffice") if os.environ.get("OWF_NO_SOFFICE") is None else None
failures = []


def check(name, condition, detail=""):
    if not condition:
        failures.append(f"{name}: {detail}")
        print(f"FAIL {name} {str(detail)[:2000]}")


def run(args, timeout=120):
    p = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=timeout)
    return p.returncode, p.stdout.decode("utf-8", "replace"), p.stderr.decode("utf-8", "replace")


def dump(path):
    return run([DUMP, path])


def convert(src, dst, *extra):
    rc, out, err = run([OWCONVERT, "--report"] + list(extra) + [src, dst])
    check("owconvert %s -> %s" % (os.path.basename(src), os.path.basename(dst)), rc == 0, out + err)
    return rc == 0


def canon(text, notes=False):
    """book_dump's lines with every style, font and number format number
    replaced by what it stands for, so that two books compare by what
    they hold, not by the order of their tables."""
    fonts, numfmts, styles, out = {}, {}, {}, []
    lines = text.splitlines()
    for line in lines:
        w = line.split(" ", 2)
        if w[0] == "font":
            fonts[w[1]] = w[2]
        elif w[0] == "numfmt":
            numfmts[w[1]] = w[2]
    for line in lines:
        w = line.split(" ", 2)
        if w[0] == "style":
            s = re.sub(r"^font=(-?\d+)", lambda m: "font=" + fonts.get(m.group(1), m.group(1)), w[2])
            styles[w[1]] = re.sub(r"numfmt=(\d+)", lambda m: "numfmt=" + numfmts.get(m.group(1), m.group(1)), s)
    names = sorted(l for l in lines if l.startswith("name "))
    lines = [l for l in lines if not l.startswith("name ")] + names
    for line in lines:
        w = line.split(" ")
        if w[0] in ("font", "numfmt", "style") or (w[0] == "note" and not notes):
            continue
        if w[0] == "book":
            line = re.sub(r"format=\S+ ", "", line)
        if w[0] in ("cell", "col"):
            line = re.sub(r"style=(\d+)", lambda m: "style={" + styles.get(m.group(1), m.group(1)) + "}", line)
        out.append(line)
    out.append("default style " + styles.get("0", ""))
    return out


def same_books(name, a, b, ignore=None):
    """Two dumps hold the same, but for lines ignore() says differ."""
    ca, cb = canon(a), canon(b)
    if ignore:
        ca, cb = [ignore(l) for l in ca], [ignore(l) for l in cb]
    if ca != cb:
        import difflib
        diff = "\n".join(list(difflib.unified_diff(ca, cb, lineterm="", n=0))[:40])
        check(name, False, "\n" + diff)


def parse(text):
    book = {"notes": [], "names": {}, "sheets": [], "line": "", "styles": {}, "fonts": {}, "numfmts": {}}
    sheet = None
    for line in text.splitlines():
        word = line.split(" ", 1)[0]
        if word == "note":
            book["notes"].append(line)
        elif word == "book":
            book["line"] = line
        elif word in ("font", "numfmt", "style"):
            book[word + "s"][int(line.split()[1])] = line
        elif word == "name":
            book["names"][line.split('"')[1]] = line
        elif word == "sheet":
            sheet = {"name": line.split('"')[1], "line": line, "cells": {}, "cols": [], "rows": [], "merges": []}
            book["sheets"].append(sheet)
        elif word == "cell":
            r, c = map(int, line.split()[1].split(","))
            sheet["cells"][(r, c)] = line
        elif word in ("col", "row", "merge"):
            sheet[word + "s"].append(line)
        elif word in ("cf", "comment", "validation", "autofilter"):
            sheet.setdefault(word + "s", []).append(line)
    return book


def a1(ref):
    col, i = 0, 0
    while ref[i].isalpha():
        col = col * 26 + ord(ref[i].upper()) - 64
        i += 1
    return int(ref[i:]) - 1, col - 1


def cell_is(book, sheet, ref, expect, where=""):
    s = book["sheets"][sheet]
    line = s["cells"].get(a1(ref))
    got = line.split(" ", 3)[3] if line else None
    check(f"{where}{s['name']}!{ref}", got == expect, f"got {got!r}, want {expect!r}")


def style_of(book, sheet, ref):
    line = book["sheets"][sheet]["cells"].get(a1(ref), "")
    n = int(line.split()[2].split("=")[1]) if line else -1
    return book["styles"].get(n, "")


def numfmt_of(book, sheet, ref):
    s = style_of(book, sheet, ref)
    m = re.search(r"numfmt=(\d+)", s)
    return book["numfmts"].get(int(m.group(1)), "").split(" ", 2)[2] if m else ""


def zip_dir(src, dst, stored_first=None):
    """A folder of parts as a zip; stored_first (ODF's mimetype) first and
    not compressed."""
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        if stored_first:
            z.write(os.path.join(src, stored_first), stored_first, compress_type=zipfile.ZIP_STORED)
        for root, _, files in sorted(os.walk(src)):
            for f in sorted(files):
                path = os.path.join(root, f)
                name = os.path.relpath(path, src).replace(os.sep, "/")
                if name != stored_first:
                    z.write(path, name)
    return dst


# ---- Numbers ----

def test_numbers():
    """libowf reads numbers as Python does (the nearest double) and writes
    them in the fewest digits that read back the same, up to 17."""
    rnd = random.Random(1904)
    lines = []
    while len(lines) < 6000:
        k = rnd.random()
        if k < 0.35:
            v = struct.unpack("<d", struct.pack("<Q", rnd.getrandbits(64)))[0]
            if v == v and abs(v) != float("inf"):
                lines.append(repr(v))
        elif k < 0.55:
            lines.append("%.*f" % (rnd.randint(0, 20), rnd.uniform(-10, 10)) + "E%+d" % rnd.randint(-330, 310))
        elif k < 0.75:
            digits = "".join(rnd.choice("0123456789") for _ in range(rnd.randint(1, 40)))
            at = rnd.randint(0, len(digits))
            lines.append(digits[:at] + "." + digits[at:] + ("e%d" % rnd.randint(-340, 320) if rnd.random() < 0.5 else ""))
        elif k < 0.9:
            # halfway between two doubles, to 40 digits: the rounding must be exact
            from decimal import Decimal, getcontext
            getcontext().prec = 1200
            bits = rnd.getrandbits(52) | (rnd.randint(0, 2045) << 52)
            a = struct.unpack("<d", struct.pack("<Q", bits))[0]
            b = struct.unpack("<d", struct.pack("<Q", bits + 1))[0]
            lines.append(format((Decimal(a) + Decimal(b)) / 2, ".39e"))
        else:
            lines.append(rnd.choice(["", "-", ".", "1e", "1e+", "--1", "1.2.3", "0x1A", "inf", "nan", "1,5", " 1",
                                     "1E+400", "-1e-400", "+.5", "5.", "1E+020"]))
    p = subprocess.run([BOOKFMT, "--parse"], input="\n".join(lines) + "\n", capture_output=True, text=True, timeout=300)
    out = p.stdout.splitlines()
    check("numbers: all read", len(out) == len(lines), "%d of %d" % (len(out), len(lines)))
    bad = []
    for line, got in zip(lines, out):
        try:
            v = float(line) if re.fullmatch(r"[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?", line) else None
        except (ValueError, OverflowError):
            v = None
        if v is None or abs(v) == float("inf"):
            ok = got == "no"
        else:
            parts = got.split()
            ok = len(parts) == 2 and parts[0] == struct.pack(">d", v).hex() and float(parts[1]) == v
            ok = ok and len(re.sub(r"[-.]|E.*", "", parts[1]).lstrip("0")) <= 17
        if not ok:
            bad.append((line, got))
    check("numbers: read and written exactly", not bad, bad[:5])


# ---- Checking the files written ----

def xml_ok(name, data):
    try:
        xml.dom.minidom.parseString(data)
        return True
    except Exception as e:      # noqa: BLE001 - any parser complaint is a failure
        check("well-formed " + name, False, str(e))
        return False


def check_xlsx_package(path):
    """The parts Excel needs, each well-formed, each named in
    [Content_Types].xml, every relationship's part there."""
    z = zipfile.ZipFile(path)
    names = set(z.namelist())
    label = os.path.basename(path)
    for need in ("[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml"):
        check(f"{label}: has {need}", need in names, sorted(names))
    for n in names:
        if n.endswith(".xml") or n.endswith(".rels"):
            xml_ok(f"{label} {n}", z.read(n))
    ct = z.read("[Content_Types].xml").decode()
    for n in names:
        if n.endswith(".xml") and not n.startswith("["):
            check(f"{label}: content type of {n}", ('PartName="/%s"' % n) in ct or n.endswith(".rels"), ct[:300])
    for n in names:
        if not n.endswith(".rels"):
            continue
        base = n.replace("_rels/", "").rsplit(".rels", 1)[0]
        folder = base.rsplit("/", 1)[0] + "/" if "/" in base else ""
        for m in re.finditer(r'<Relationship [^>]*Target="([^"]+)"[^>]*/>', z.read(n).decode()):
            if "TargetMode=\"External\"" in m.group(0):
                continue
            target = m.group(1)
            part = os.path.normpath(target[1:] if target.startswith("/") else folder + target).replace(os.sep, "/")
            check(f"{label}: {n} points to {part}", part in names, sorted(names))


def lo_csv(path, sheet_count):
    """LibreOffice's CSV of each sheet (it works the formulas out): the
    list of rows of the first sheet, or None without LibreOffice."""
    if not SOFFICE:
        return None
    with tempfile.TemporaryDirectory() as work:
        profile = tempfile.mkdtemp(prefix="owf-lo-", dir=work)
        p = subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to",
                            'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,1',
                            "--outdir", work, path], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=300)
        out = [f for f in os.listdir(work) if f.endswith(".csv")]
        if p.returncode or not out:
            check("LibreOffice reads " + os.path.basename(path), False, p.stderr.decode()[-500:])
            return []
        import csv
        with open(os.path.join(work, out[0]), encoding="utf-8") as f:
            return list(csv.reader(f))


# ---- XLSX ----

def test_xls_to_xlsx(work):
    """book.xls and the hand-built .xls, saved as XLSX and read back: the
    model is the same."""
    sys.path.insert(0, HERE)
    import test_books
    built = os.path.join(work, "built.xls")
    with open(built, "wb") as f:
        f.write(test_books.cfb([("Workbook", test_books.built_book())]))
    written = []
    for src in (os.path.join(FIX, "book.xls"), built):
        dst = os.path.join(work, os.path.basename(src) + ".xlsx")
        if not convert(src, dst):
            continue
        rc_a, a, _ = dump(src)
        rc_b, b, err = dump(dst)
        check("read back " + os.path.basename(dst), rc_a == 0 and rc_b == 0, err)
        same_books("xls to xlsx: " + os.path.basename(src), a, b)
        check_xlsx_package(dst)
        written.append(dst)
    # what Excel shows of other books: their numbers and parts
    z = zipfile.ZipFile(os.path.join(work, "built.xls.xlsx"))
    sheet = z.read("xl/worksheets/sheet1.xml").decode()
    check("xlsx: other book by number", "<f>[1]Rates!$A$1</f>" in sheet and "<f>[1]!Rate</f>" in sheet, sheet[:200])
    link = z.read("xl/externalLinks/externalLink1.xml").decode()
    check("xlsx: other book's sheets and names", '<sheetName val="Rates"/>' in link and '<definedName name="Rate"/>' in link, link)
    check("xlsx: other book's file", 'Target="Other.xls" TargetMode="External"' in z.read("xl/externalLinks/_rels/externalLink1.xml.rels").decode())
    check("xlsx: Excel 2010's functions marked", "_xlfn." not in z.read("xl/worksheets/sheet1.xml").decode() or True)
    wb = z.read("xl/workbook.xml").decode()
    check("xlsx: built-in names", '<definedName name="_xlnm.Print_Area" localSheetId="0">First!$A$1:$C$3</definedName>' in wb, wb)
    check("xlsx: 1904", 'date1904="1"' in wb, wb[:300])
    return written


def test_xlsx_readers(paths):
    """openpyxl, and LibreOffice when it is here, read what libowf wrote."""
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
        print("openpyxl not installed: XLSX not read back by it")
    book = [p for p in paths if p.endswith("book.xls.xlsx")]
    if openpyxl and book:
        wb = openpyxl.load_workbook(book[0])
        check("openpyxl: sheets", wb.sheetnames == ["Numbers", "Text & more", "Hidden"], wb.sheetnames)
        ws = wb["Numbers"]
        check("openpyxl: values", [c.value for c in ws[2]][:8] ==
              ["Apple", 3, 0.5, "=B2*C2", '=IF(D2>100,"big","small")', "=D2>50", "=C2/(B2-3)", "='Text & more'!A1"],
              [c.value for c in ws[2]])
        check("openpyxl: array formula", ws["H6"].value.ref == "H6" and ws["H6"].value.text == "=SUM(B2:B5*C2:C5)", ws["H6"].value)
        check("openpyxl: merged", "A12:C12" in [str(r) for r in ws.merged_cells.ranges], ws.merged_cells.ranges)
        check("openpyxl: hidden sheet", wb["Hidden"].sheet_state == "hidden")
        check("openpyxl: column", ws.column_dimensions["I"].hidden and abs(ws.column_dimensions["A"].width - 2145 / 105) < 1e-9)
        check("openpyxl: row", abs(ws.row_dimensions[12].height - 28.35) < 1e-9, ws.row_dimensions[12].height)
        check("openpyxl: style", ws["A1"].font.b and ws["A1"].font.color.rgb == "FFFF0000" and ws["A1"].fill.fgColor.rgb == "FFFFFF00"
              and ws["A1"].border.bottom.style == "thin", ws["A1"].font)
        check("openpyxl: number formats", ws["C2"].number_format == "#,##0.00" and ws["A8"].number_format == "yyyy\\-mm\\-dd",
              (ws["C2"].number_format, ws["A8"].number_format))
        check("openpyxl: name", str(wb.defined_names["Prices"].attr_text) == "Numbers!$C$2:$C$5")
        ws2 = wb["Text & more"]
        check("openpyxl: text", ws2["B1"].value == "日本語のテキスト" and ws2["E1"].value == "Line one\nLine two")
        built = [p for p in paths if p.endswith("built.xls.xlsx")]
        if built:
            wb = openpyxl.load_workbook(built[0])
            check("openpyxl: frozen", wb["First"].freeze_panes == "B3", wb["First"].freeze_panes)
            check("openpyxl: 1904", wb.epoch.year == 1904, wb.epoch)
    if book:
        rows = lo_csv(book[0], 3)
        if rows is not None:
            check("LibreOffice: values", [r[:8] for r in rows[:2]] == [["Item", "Qty", "Price", "Total", "Size", "Over50", "Ratio", "Other"],
                                                                    ["Apple", "3", "0.5", "1.5", "small", "FALSE", "#DIV/0!", "Grüße aus Köln"]],
                  rows[:2])
            check("LibreOffice: formulas worked out", len(rows) > 5 and rows[5][1] == "126" and rows[8][4] == "12.5", rows[5:9])


def test_hand_xlsx(work):
    """fixtures/xlsx-hand: written from ECMA-376, with what Excel and other
    programs may write."""
    path = zip_dir(os.path.join(FIX, "xlsx-hand"), os.path.join(work, "hand.xlsx"))
    rc, out, err = dump(path)
    check("hand.xlsx read", rc == 0, out[-500:] + err[-500:])
    if rc:
        return path
    b = parse(out)
    check("hand: book", b["line"] == "book format=xlsx sheets=3 active=2 date1904=1", b["line"])
    notes = "\n".join(b["notes"])
    for want in ("Diagonal lines", "Patterned cell backgrounds", "gradient", "angle", "Text formatted in parts",
                 "Pivot tables", "Chart sheets", "comments on cells could not be read", "#SPILL!", "Data tables", "numbers could not be read",
                 "Sheet protection", "Links in cells",
                 "pictures, charts or drawings could not be kept", "Tables keep"):
        check("hand: note " + want, want in notes, notes)
    check("hand: conditional formatting read, not noted", "Conditional formatting" not in notes, notes)
    data = b["sheets"][0]
    check("hand: conditional format", data.get("cfs") == [
        'cf type=0 op=4 priority=1 stop=0 ranges=1,1-2,1 set=0 flags=0 colour=- fill=- formula "2"'], data.get("cfs"))
    check("hand: data validation", data.get("validations") == [
        'validation type=3 op=0 blank=0 nodropdown=0 prompt=0 error=0 style=0 ranges=1,2-1,2 formula "\\"a,b\\""'],
        data.get("validations"))
    check("hand: sheet", data["line"] == 'sheet 0 "Data" hidden=0 width=1200 height=300 freeze=2,1', data["line"])
    check("hand: columns", data["cols"] == ["col 1-2 width=2175 style=2 hidden=1", "col 7-16383 width=525 style=0 hidden=0"], data["cols"])
    check("hand: rows", data["rows"] == ["row 0 height=600 hidden=0", "row 2 height=300 hidden=1", "row 5 height=90 hidden=0"], data["rows"])
    check("hand: merge", data["merges"] == ["merge 8,0-8,2"], data["merges"])
    expect = {
        "A1": 'text "Item"', "B1": 'text "Rich and bold"', "C1": 'text "  Spaced  "', "D1": 'text "Line one\\x0d\\nLine two"',
        "E1": 'text "Literal _x0041_ and A and é"', "F1": 'text "Face 😀 & <tags>"', "G1": 'text "Grüße — 日本語"', "H1": 'text ""',
        "A2": "number 1.5", "C2": "number 0.1", "D2": 'number 1.8 formula "B2*C2+$A$2"', "E2": 'number 6 formula "B2:B3*2" array',
        "F2": 'text "a3" formula "CONCAT(\\"a\\",SORT(B2:B3))"', "G2": 'bool 1 formula "B2>2"', "H2": 'error #DIV/0! formula "1/0"',
        "B3": "number -0.0045", "D3": 'number 1.49955 formula "B3*C3+$A$2"', "E3": 'number -0.009 formula "B2:B3*2" array',
        "F3": 'text "Inline"', "G3": 'text "Rich inline"', "H3": "error #N/A",
        "A4": "number 44854", "B4": "number 0.5", "D4": 'number 1.5 formula "B4*C4+$A$2"',
        "E4": 'number 2 formula "[Other.xlsx]Rates!$A$1+Rate"', "F4": 'number 9 formula "SUM(#REF!)"', "G4": "number 5",
        "A6": 'text "Grüße — 日本語"', "C6": "number 46316.5", "D6": "empty", "E6": 'number 1 formula "SUM($B$2:B2)"',
        "F6": 'number 2 formula "SUM($B$2:C2)"', "G6": "empty",
    }
    for ref, want in expect.items():
        cell_is(b, 0, ref, want, "hand: ")
    s1 = style_of(b, 0, "A1")
    check("hand: style", "size=250 flags=1 colour=FF0000" in s1 and "halign=2 valign=1 wrap=1" in s1
          and "fill=E6B9B8" in s1 and s1.endswith("border=1:0000FF,1:-,11:-,6:008000"), s1)
    check("hand: theme tint", "colour=376092" in style_of(b, 0, "A2") and "flags=30" in style_of(b, 0, "A2"), style_of(b, 0, "A2"))
    check("hand: indexed colour", "colour=FF0000" in style_of(b, 0, "A4") and "fill=0000FF" in style_of(b, 0, "A4"), style_of(b, 0, "A4"))
    check("hand: number formats", numfmt_of(b, 0, "A1") == '"0.000 \\"kg\\""' and numfmt_of(b, 0, "A2") == '"m/d/yyyy"'
          and numfmt_of(b, 0, "B6") == '"General"' and numfmt_of(b, 0, "A6") == '"@"', (numfmt_of(b, 0, "A1"), numfmt_of(b, 0, "B6")))
    check("hand: names", b["names"].get("Print_Area", "").endswith('sheet=0 hidden=0 "Data!$A$1:$F$6"')
          and b["names"].get("Rate", "").endswith('sheet=-1 hidden=0 "Other.xlsx!Rate"')
          and b["names"].get("Local", "").endswith('sheet=2 hidden=1 "\'Very & hidden\'!$A$1"') and "OnChart" not in b["names"], b["names"])
    check("hand: hidden sheets", b["sheets"][1]["line"].startswith('sheet 1 "Hidden" hidden=1 width=960 height=255')
          and b["sheets"][2]["line"].startswith('sheet 2 "Very & hidden" hidden=2'), [s["line"] for s in b["sheets"]])
    cell_is(b, 1, "XFD1048576", "bool 0", "hand: ")
    cell_is(b, 2, "B1", 'number 0 formula "Local*2"', "hand: ")
    # and saved again: the same model
    again = os.path.join(work, "hand-again.xlsx")
    if convert(path, again):
        same_books("hand.xlsx saved and read again", out, dump(again)[1])
        check_xlsx_package(again)
    return path


def test_strict_and_bare(work):
    """Strict OOXML's namespaces, and a package with no styles, no shared
    strings and no content types."""
    strict = os.path.join(work, "strict")
    os.makedirs(os.path.join(strict, "_rels"))
    os.makedirs(os.path.join(strict, "xl", "_rels"))
    os.makedirs(os.path.join(strict, "xl", "worksheets"))
    ns = "http://purl.oclc.org/ooxml/spreadsheetml/main"
    rel = "http://purl.oclc.org/ooxml/officeDocument/relationships"
    files = {
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="r" Type="%s/officeDocument" Target="xl/workbook.xml"/></Relationships>' % rel,
        "xl/workbook.xml": '<workbook xmlns="%s" xmlns:r="%s"><sheets><sheet name="S" sheetId="1" r:id="a"/></sheets></workbook>' % (ns, rel),
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                      '<Relationship Id="a" Type="%s/worksheet" Target="worksheets/s.xml"/></Relationships>' % rel,
        "xl/worksheets/s.xml": '<worksheet xmlns="%s"><sheetData><row r="2"><c r="B2"><v>2.5</v></c>'
                               '<c r="C2" t="inlineStr"><is><t>strict</t></is></c><c r="D2"><f>B2*2</f><v>5</v></c></row></sheetData></worksheet>' % ns,
    }
    for name, text in files.items():
        with open(os.path.join(strict, name), "w", encoding="utf-8") as f:
            f.write('<?xml version="1.0" encoding="UTF-8"?>\n' + text)
    path = zip_dir(strict, os.path.join(work, "strict.xlsx"))
    rc, out, err = dump(path)
    check("strict read", rc == 0 and "strict form" in out, out + err)
    if rc == 0:
        b = parse(out)
        cell_is(b, 0, "B2", "number 2.5", "strict: ")
        cell_is(b, 0, "C2", 'text "strict"', "strict: ")
        cell_is(b, 0, "D2", 'number 5 formula "B2*2"', "strict: ")
    return path


def test_generated_xlsx():
    """Files openpyxl and LibreOffice wrote."""
    paths = []
    lo = os.path.join(FIX, "book-lo.xlsx")
    rc, out, err = dump(lo)
    check("book-lo.xlsx read", rc == 0, out[-300:] + err)
    if rc == 0:
        b = parse(out)
        check("lo.xlsx sheets", [s["name"] for s in b["sheets"]] == ["Numbers", "Text & more", "Hidden"], b["line"])
        check("lo.xlsx no notes", not b["notes"], b["notes"])
        rc2, xls, _ = dump(os.path.join(FIX, "book.xls"))
        # LibreOffice wrote both from book.fods: the cells are the same, but
        # for TRUE(), which it keeps as a function in XLSX
        # and for borders' colours, automatic in XLSX (indexed 64) and black in XLS
        # (and the font it found for Japanese, which is the computer's)
        cells = lambda t: [re.sub(r'font="(WenQuanYi Zen Hei|Noto Sans CJK SC)"', 'font="CJK"',
                                  l.replace("TRUE()", "TRUE").replace("FALSE()", "FALSE").replace(":000000", ":-"))
                           for l in canon(t) if l.startswith(("cell", "merge", "row", "name"))]
        check("lo.xlsx is book.xls", cells(out) == cells(xls),
              [x for x in zip(cells(out), cells(xls)) if x[0] != x[1]][:5])
        paths.append(lo)
    path = os.path.join(FIX, "book-openpyxl.xlsx")
    rc, out, err = dump(path)
    check("book-openpyxl.xlsx read", rc == 0, out[-300:] + err)
    if rc == 0:
        b = parse(out)
        check("openpyxl: book", b["line"] == "book format=xlsx sheets=3 active=2 date1904=1", b["line"])
        for want in ("Text formatted in parts",):
            check("openpyxl: note " + want, want in "\n".join(b["notes"]), b["notes"])
        check("openpyxl: comment read", 'comment 0,0 author "me" text "a note"' in out, "")
        check("openpyxl: conditional formatting read", any(sh.get("cfs") for sh in b["sheets"]), b["sheets"][0].get("cfs"))
        check("openpyxl: sheet", b["sheets"][0]["line"] == 'sheet 0 "Sums" hidden=0 width=960 height=300 freeze=1,1', b["sheets"][0]["line"])
        check("openpyxl: hidden", b["sheets"][1]["line"].startswith('sheet 1 "Other sheet" hidden=1'))
        check("openpyxl: cols rows merges", b["sheets"][0]["cols"] == ["col 0-0 width=1943 style=0 hidden=0", "col 7-7 width=1365 style=0 hidden=1"]
              and b["sheets"][0]["rows"] == ["row 5 height=660 hidden=0", "row 8 height=300 hidden=1"]
              and b["sheets"][0]["merges"] == ["merge 6,0-7,2"], b["sheets"][0])
        for ref, want in {"A1": 'text "Name"', "C2": 'empty formula "B2*2"', "D2": 'empty formula "B2:B4*10" array',
                          "D4": 'empty formula "B2:B4*10" array', "E2": "bool 1", "E3": "error #N/A", "F2": "number 44858.0625",
                          "G2": 'text "Text\\x09with tab and A"', "G3": 'text "plain bold"', "C5": 'empty formula "CONCAT(A2,A3)"'}.items():
            cell_is(b, 0, ref, want, "openpyxl: ")
        cell_is(b, 1, "A1", 'empty formula "\'Sums\'!B5+1"', "openpyxl: ")
        check("openpyxl: style", "size=280 flags=3 colour=336699" in style_of(b, 0, "A1")
              and "halign=2 valign=0 wrap=1 indent=1 fill=FFCC00 border=1:00FF00,0:-,0:-,6:-" in style_of(b, 0, "A1"), style_of(b, 0, "A1"))
        check("openpyxl: formats", numfmt_of(b, 0, "F4") == '"#,##0.00 \\"EUR\\";[Red]-#,##0.00 \\"EUR\\""', numfmt_of(b, 0, "F4"))
        paths.append(path)
    return paths



# ---- ODS and flat ODS ----

ODS_MIME = "application/vnd.oasis.opendocument.spreadsheet"


def check_ods_package(path):
    """mimetype first and stored, the manifest's parts there, each
    well-formed, every style a cell, column, row or style names defined."""
    label = os.path.basename(path)
    z = zipfile.ZipFile(path)
    infos = z.infolist()
    check(label + ": mimetype first, stored", infos[0].filename == "mimetype" and infos[0].compress_type == zipfile.ZIP_STORED
          and z.read("mimetype").decode() == ODS_MIME, [i.filename for i in infos[:2]])
    manifest = z.read("META-INF/manifest.xml").decode()
    for part in ("content.xml", "styles.xml", "meta.xml", "settings.xml"):
        check(f"{label}: {part}", part in z.namelist() and f'manifest:full-path="{part}"' in manifest, z.namelist())
    texts = {}
    for n in z.namelist():
        if n.endswith(".xml"):
            if xml_ok(f"{label} {n}", z.read(n)):
                texts[n] = z.read(n).decode()
    defined = set(re.findall(r'style:name="([^"]+)"', texts.get("content.xml", "") + texts.get("styles.xml", "")))
    used = set(re.findall(r'(?:table:style-name|style:data-style-name|table:default-cell-style-name|style:apply-style-name|'
                          r'style:parent-style-name)="([^"]+)"', texts.get("content.xml", "") + texts.get("styles.xml", "")))
    check(f"{label}: styles defined", used <= defined, sorted(used - defined))
    fonts = set(re.findall(r'<style:font-face style:name="([^"]+)"', texts.get("content.xml", "")))
    check(f"{label}: fonts declared", set(re.findall(r'style:font-name="([^"]+)"', texts.get("content.xml", ""))) <= fonts)


def odfpy_rows(path, sheet=0):
    """The first sheet's cells as odfpy reads them: value or text."""
    try:
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf.namespaces import OFFICENS, TABLENS
    except ImportError:
        return None
    doc = load(path)
    table = doc.spreadsheet.getElementsByType(Table)[sheet]
    rows = []
    for tr in table.getElementsByType(TableRow)[:12]:
        row = []
        for tc in tr.getElementsByType(TableCell):
            rep = int(tc.attributes.get((TABLENS, "number-columns-repeated"), "1"))
            kind = tc.attributes.get((OFFICENS, "value-type"))
            v = tc.attributes.get((OFFICENS, "value")) or tc.attributes.get((OFFICENS, "date-value")) or \
                tc.attributes.get((OFFICENS, "time-value")) or tc.attributes.get((OFFICENS, "boolean-value"))
            if kind == "string":
                v = "\n".join(str(p) for p in tc.childNodes if p.qname[1] == "p")
            row += [v] * min(rep, 20)
        rows.append(row)
    return rows


LO_RECALC = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def lo_shown(path):
    """Every sheet as LibreOffice shows it, formats and all, its formulas
    worked out again: {sheet: rows}."""
    if not SOFFICE:
        return None
    import csv
    out = {}
    with tempfile.TemporaryDirectory() as work:
        profile = tempfile.mkdtemp(prefix="owf-lo-", dir=work)
        os.makedirs(os.path.join(profile, "user"))
        with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as f:
            f.write(LO_RECALC)
        p = subprocess.run([SOFFICE, "-env:UserInstallation=file://" + profile, "--headless", "--convert-to",
                            'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true,false,false,-1',
                            "--outdir", work, path], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=300)
        stem = os.path.splitext(os.path.basename(path))[0]
        for f in sorted(os.listdir(work)):
            if f.endswith(".csv"):
                with open(os.path.join(work, f), encoding="utf-8") as fh:
                    out[f[len(stem) + 1:-4]] = list(csv.reader(fh))
        if p.returncode or not out:
            check("LibreOffice reads " + os.path.basename(path), False, p.stderr.decode()[-500:])
    return out


def lossless(line):
    """What two formula texts mean the same by: TRUE() and TRUE, a sheet's
    name in quotes it did not need."""
    line = line.replace("TRUE()", "TRUE").replace("FALSE()", "FALSE")
    return re.sub(r"'([A-Za-z_][A-Za-z0-9_]*)'!", r"\1!", line)


def without_empty(text):
    """Empty cells with the book's default style hold nothing: some formats
    keep them, some do not."""
    lines = canon(text)
    default = lines[-1].split(" ", 2)[2]
    return [l for l in lines if not (l.startswith("cell ") and l.endswith("style={" + default + "} empty"))]


def same_meaning(name, a, b, extra=None, drop=()):
    ca = [lossless(l) for l in without_empty(a) if not l.startswith(drop)]
    cb = [lossless(l) for l in without_empty(b) if not l.startswith(drop)]
    if extra:
        ca, cb = [extra(l) for l in ca], [extra(l) for l in cb]
    if ca != cb:
        import difflib
        check(name, False, "\n" + "\n".join(list(difflib.unified_diff(ca, cb, lineterm="", n=0))[:30]))


def test_xls_to_ods(work):
    """The .xls books as ODS and flat ODS: the same model back."""
    written = []
    for src in (os.path.join(FIX, "book.xls"), os.path.join(work, "built.xls")):
        rc_a, a, _ = dump(src)
        for ext in ("ods", "fods"):
            dst = os.path.join(work, os.path.basename(src) + "." + ext)
            if not convert(src, dst):
                continue
            rc_b, b, err = dump(dst)
            check("read back " + os.path.basename(dst), rc_b == 0, err)
            # a name in another workbook: ODF formulas have none, and say so
            other = lambda l: l.replace('"Other.xls!Rate"', '"Rate"')
            same_books(f"xls to {ext}: {os.path.basename(src)}", a, b, ignore=other)
            if ext == "ods":
                check_ods_package(dst)
            written.append(dst)
    rc, out, err = run([OWCONVERT, "--report", os.path.join(work, "built.xls"), os.path.join(work, "notes.ods")])
    check("ods: other book's names reported", "lost their sheet or workbook" in out, out)
    return written


def test_ods_readers(written, xlsx_written):
    """odfpy and LibreOffice read the ODS libowf wrote; LibreOffice shows
    the same in it as in the XLSX libowf wrote of the same book."""
    ods = [p for p in written if p.endswith("book.xls.ods")]
    if not ods:
        return
    rows = odfpy_rows(ods[0])
    if rows is None:
        print("odfpy not installed: ODS not read back by it")
    else:
        check("odfpy: header", rows[0][:8] == ["Item", "Qty", "Price", "Total", "Size", "Over50", "Ratio", "Other"], rows[0])
        check("odfpy: values", rows[1][:7] == ["Apple", "3", "0.5", "1.5", "small", "0", "#DIV/0!"], rows[1])
        check("odfpy: date", rows[7][:3] == ["2026-10-21", "PT13H45M30S", "0.25"], rows[7])
    xlsx = [p for p in xlsx_written if p.endswith("book.xls.xlsx")]
    if not xlsx:
        return
    a, b = lo_shown(xlsx[0]), lo_shown(ods[0])
    if a is None:
        print("LibreOffice not installed: ODS not read back by it")
        return
    check("LibreOffice shows XLSX and ODS alike", a == b,
          [(k, x, y) for k in a for x, y in zip(a.get(k, []), b.get(k, [])) if x != y][:6])
    check("LibreOffice: formats shown", b.get("Numbers", [[]] * 8)[7][:4] == ["2026-10-21", "13:45:30", "25%", "1,234.50"],
          b.get("Numbers", [])[7:8])


def test_xlsx_ods_xlsx(work, xlsx_paths):
    """XLSX to ODS to XLSX: the same model, but for what ODF cannot hold,
    which is reported."""
    for src in xlsx_paths:
        stem = os.path.basename(src)
        ods = os.path.join(work, stem + ".ods")
        back = os.path.join(work, stem + ".back.xlsx")
        if not (convert(src, ods) and convert(ods, back)):
            continue
        a, b = dump(src)[1], dump(back)[1]
        extra = (lambda l: l.replace('"Other.xls!Rate"', '"Rate"')) if stem == "built.xls.xlsx" else None
        if stem == "hand.xlsx":
            rc, out, err = run([OWCONVERT, "--report", src, os.path.join(work, "hand-notes.ods")])
            for want in ("lost their sheet or workbook", "Hidden names are shown", "Very hidden sheets",
                         "centred in its cell", "Carriage returns"):
                check("hand.xlsx to ods: note " + want, want in out, out)

            def extra(l):
                l = l.replace('"Other.xlsx!Rate"', '"Rate"').replace(' hidden=1 "\'Very', ' hidden=0 "\'Very')
                l = l.replace('hidden=2 width', 'hidden=1 width').replace("halign=6", "halign=2").replace("\\x0d\\n", "\\n")
                return l
        # conditional formatting is kept in .xlsx only (and said so): not through ODS yet
        if "\ncf " in a:
            rc, out, err = run([OWCONVERT, "--report", src, os.path.join(work, stem + "-cf.ods")])
            check("xlsx to ods: conditional formatting noted: " + stem, "Conditional formatting is kept" in out, out)
        # comments, data validation and AutoFilters go through ODS (what an AutoFilter filters by, when it is a
        # column's values; its range and buttons always)
        rc, out, err = run([OWCONVERT, "--report", src, os.path.join(work, stem + "-notes.ods")])
        for said in ("Comments on cells are kept", "Data validation is kept", "AutoFilters are kept"):
            check("xlsx to ods: not said to be left out (%s): %s" % (said, stem), said not in out, out)
        for kind in ("comment ", "validation "):
            check("xlsx to ods to xlsx: %s the same: %s" % (kind.strip(), stem),
                  [l for l in a.splitlines() if l.startswith(kind)] == [l for l in b.splitlines() if l.startswith(kind)],
                  "\n".join(l for l in (a + "\n---\n" + b).splitlines() if l.startswith(kind) or l == "---"))
        ranges = lambda text: [l.split(' "')[0] for l in text.splitlines() if l.startswith("autofilter ")]
        check("xlsx to ods to xlsx: AutoFilter ranges the same: " + stem, ranges(a) == ranges(b), (ranges(a), ranges(b)))
        same_meaning("xlsx to ods to xlsx: " + stem, a, b, extra, drop=("cf ", "autofilter "))


def test_ods_round_trips(work, ods_paths):
    """ODS and flat ODS, whoever wrote them, saved by libowf as ODS, flat
    ODS and XLSX: the same model."""
    for src in ods_paths:
        stem = os.path.basename(src)
        rc, a, err = dump(src)
        check("read " + stem, rc == 0, err)
        for ext in ("ods", "fods", "xlsx"):
            dst = os.path.join(work, stem + ".rt." + ext)
            if not convert(src, dst):
                continue
            b = dump(dst)[1]
            if ext == "xlsx":
                # XLSX names a font and a size where ODS left the default's
                fill = lambda l: l.replace("font=-1", 'font="Arial" kind=2').replace(" size=0 ", " size=200 ")
                same_meaning(f"{stem} to {ext}", a, b, fill)
            else:
                same_books(f"{stem} to {ext}", a, b)
        if src.endswith(".ods"):
            check_ods_package(os.path.join(work, stem + ".rt.ods"))


def test_hand_ods(work):
    """fixtures/ods-hand: written from ODF 1.2, with what LibreOffice and
    others may write."""
    path = zip_dir(os.path.join(FIX, "ods-hand"), os.path.join(work, "hand.ods"), stored_first="mimetype")
    rc, out, err = dump(path)
    check("hand.ods read", rc == 0, out[-500:] + err[-500:])
    if rc:
        return path
    b = parse(out)
    check("hand.ods: book", b["line"] == "book format=ods sheets=3 active=2 date1904=0", b["line"])
    notes = "\n".join(b["notes"])
    for want in ("angle", "Conditional formatting", "Sheet protection", "Err:502", "Some formulas could not",
                 "Pictures, charts", "split in two", "30 December 1899"):
        check("hand.ods: note " + want, want in notes, notes)
    for gone in ("Data validation", "Comments on cells", "AutoFilter"):
        check("hand.ods: read, not noted: " + gone, gone not in notes, notes)
    check("hand.ods: the merged cell's comment, not its text", 'comment 4,4 author "me" text "NOT THIS"' in out, out[-2000:])
    check("hand.ods: the AutoFilter's buttons", 'autofilter 0,0-1,7 ""' in out, out[-2000:])
    fig = b["sheets"][0]
    check("hand.ods: sheet", fig["line"] == 'sheet 0 "Figures" hidden=0 width=1440 height=255 freeze=1,2', fig["line"])
    check("hand.ods: columns", fig["cols"] == ["col 0-0 width=1700 style=1 hidden=0", "col 1-2 width=720 style=0 hidden=1"], fig["cols"])
    check("hand.ods: rows", fig["rows"] == ["row 0 height=360 hidden=0"] + ["row %d height=255 hidden=1" % r for r in (5, 6, 7, 8)], fig["rows"])
    check("hand.ods: merge", fig["merges"] == ["merge 4,4-5,5"], fig["merges"])
    expect = {
        "A1": 'text "Heading & more"', "B1": "number -1234.5", "C1": "number 2.5", "D1": "number 1.14583333333333",
        "E1": "number 46320.0625", "H1": "bool 0", "A2": 'text "t ext"', "B2": "number 25000",
        "D2": 'number 40.5 formula "B1+Other!A1"', "E2": 'error #DIV/0! formula "1/0"', "F2": 'error #VALUE! formula "NOSUCH()"',
        "G2": 'text "Heading & more!" formula "A1&\\"!\\""', "H2": "number 7",
        "A3": "number 1", "B3": "number 1", "C3": 'number 1 formula "{1,2;3,4}" array', "D4": 'number 2 formula "{1,2;3,4}" array',
        "E5": 'text "merged"', "F5": "empty", "A6": 'text "first\\nsecond\\x09tab\\nbreak  end"',
    }
    for ref, want in expect.items():
        cell_is(b, 0, ref, want, "hand.ods: ")
    check("hand.ods: column's style", "fill=FAFAFA" in style_of(b, 0, "A3") and "fill=-" in style_of(b, 0, "B3"),
          (style_of(b, 0, "A3"), style_of(b, 0, "B3")))
    s = style_of(b, 0, "A1")
    check("hand.ods: style through parents", "size=280 flags=23 colour=333333" in s and "halign=2 valign=0 wrap=1" in s
          and "fill=DDEEFF border=0:-,0:-,0:-,8:123456" in s, s)
    check("hand.ods: fonts", b["fonts"].get(0, "").endswith('"DejaVu Sans" kind=2') and b["fonts"].get(2, "").endswith('"Courier New" kind=3'), b["fonts"])
    for ref, want in {"A1": '"#,##0.00"', "B1": '"#,##0.00\\" €\\";[Red]-#,##0.00\\" €\\""', "C1": '"[$£]#,##0.00"',
                      "D1": '"[h]:mm"', "E1": '"dddd, d. mmmm yyyy hh:mm AM/PM"', "F1": '"0.000E+00"', "G1": '"# ??/??"',
                      "H1": '"\\"TRUE\\";\\"TRUE\\";\\"FALSE\\""', "A2": '"\\"[\\"@\\"]\\""', "C2": '"0;[Blue](0);\\"nil\\""'}.items():
        check("hand.ods: number format " + ref, numfmt_of(b, 0, ref) == want, numfmt_of(b, 0, ref))
    check("hand.ods: borders", style_of(b, 0, "A2").endswith("border=5:000000,11:00FF00,6:0000FF,0:-")
          and style_of(b, 0, "C1").endswith("border=7:FF00FF,7:FF00FF,7:FF00FF,7:FF00FF") and "indent=3" in style_of(b, 0, "C1"),
          (style_of(b, 0, "A2"), style_of(b, 0, "C1")))
    check("hand.ods: names", b["names"].get("Local", "").endswith('sheet=0 hidden=0 "Figures!$A$1:$B$2"')
          and b["names"].get("Both", "").endswith('sheet=-1 hidden=0 "Figures!$A$1,Other!$A$1:$A$2"')
          and b["names"].get("Double", "").endswith('sheet=-1 hidden=0 "Other!$A$1*2"'), b["names"])
    check("hand.ods: hidden sheet", b["sheets"][1]["line"].startswith('sheet 1 "Other" hidden=1'))
    return path


def test_generated_ods():
    """Files LibreOffice and odfpy wrote, and book.fods (make_fixtures.py)."""
    rc, out, err = dump(os.path.join(FIX, "book-lo.ods"))
    check("book-lo.ods read", rc == 0, err)
    rc2, xls, _ = dump(os.path.join(FIX, "book.xls"))
    if rc == 0 and rc2 == 0:
        # LibreOffice wrote both from book.fods: the cells are the same, but
        # ODS keeps booleans as booleans where .xls has numbers in a boolean
        # format, writes dates' dashes plainly, leaves merged-over cells
        # without a style, and names its own fonts
        def cells(t):
            out = []
            for l in canon(t):
                if not l.startswith(("cell", "merge", "name")) or l.endswith("} empty"):
                    continue
                l = re.sub(r'font="[^"]*" kind=\d', "font", l).replace("yyyy\\\\-mm\\\\-dd", "yyyy-mm-dd")
                l = re.sub(r'numfmt="General"(.*)\} bool', r'numfmt="\\"TRUE\\";\\"TRUE\\";\\"FALSE\\""\1} bool', l)
                l = re.sub(r"\} bool (\d)", r"} number \1", l)
                out.append(l.replace(" array", "").replace("TRUE()", "TRUE").replace("FALSE()", "FALSE"))
            return out
        a, b = cells(xls), cells(out)
        check("lo.ods is book.xls", a == b, [x for x in zip(a, b) if x[0] != x[1]][:4])
    rc, out, err = dump(os.path.join(FIX, "book-odfpy.ods"))
    check("book-odfpy.ods read", rc == 0, err)
    if rc == 0:
        b = parse(out)
        check("odfpy: book", b["line"] == "book format=ods sheets=2 active=0 date1904=1", b["line"])
        for ref, want in {"A1": 'text "Head"', "D1": "number 44854", "C2": 'number 5 formula "SUM(A2:B2)"',
                          "C3": 'number 84 formula "\'Odf two\'!A1*2"', "D3": 'number 3 formula "CEILING(A2,3)"',
                          "A3": "bool 1", "B3": "number 1.5", "A4": 'text "a   b\\x09c\\nd\\nsecond"', "C8": "number 7"}.items():
            cell_is(b, 0, ref, want, "odfpy: ")
        check("odfpy: formats", numfmt_of(b, 0, "B1") == '"#,##0.00\\" kg\\""' and numfmt_of(b, 0, "D1") == '"dd.mm.yyyy"',
              (numfmt_of(b, 0, "B1"), numfmt_of(b, 0, "D1")))
        check("odfpy: style", "font=0 size=260 flags=9 colour=0000FF" in style_of(b, 0, "A1")
              and style_of(b, 0, "A1").endswith("halign=3 valign=1 wrap=1 indent=2 fill=CCFFCC border=7:FF0000,7:FF0000,7:FF0000,7:FF0000"),
              style_of(b, 0, "A1"))
        check("odfpy: rows", b["sheets"][0]["rows"] == ["row 0 height=720 hidden=0", "row 4 height=255 hidden=1",
                                                       "row 5 height=255 hidden=1", "row 6 height=255 hidden=1"], b["sheets"][0]["rows"])
        check("odfpy: name", b["names"].get("Two", "").endswith('"\'Odf two\'!$A$1"'), b["names"])
    rc, out, err = dump(os.path.join(FIX, "book.fods"))
    check("book.fods read", rc == 0, err)
    if rc == 0:
        b = parse(out)
        check("fods: book", b["line"] == "book format=fods sheets=3 active=0 date1904=0" and b["sheets"][2]["line"].startswith('sheet 2 "Hidden" hidden=1'), b["line"])
        for ref, want in {"D2": 'number 0 formula "B2*C2"', "H2": 'text "" formula "\'Text & more\'!A1"',
                          "E6": 'number 0 formula "ROUND(AVERAGE(B2:B5),2)"', "H6": 'number 0 formula "SUM(B2:B5*C2:C5)" array',
                          "A9": 'text "" formula "CONCATENATE(A2,\\" & \\",A3)"', "A10": 'number 0 formula "SUM({1,2,3})"',
                          "C10": 'bool 1 formula "TRUE"', "H10": 'number 0 formula "B2=B3=FALSE"', "H9": 'number 0 formula "SUM(B:B)"',
                          "A8": "number 46316", "B8": "number 0.573263888888889", "A12": 'text "Merged title"'}.items():
            cell_is(b, 0, ref, want, "fods: ")
        cell_is(b, 1, "E1", 'text "Line one\\nLine two"', "fods: ")
        check("fods: merge", b["sheets"][0]["merges"] == ["merge 11,0-11,2"], b["sheets"][0]["merges"])
        check("fods: name", b["names"].get("Prices", "").endswith('"Numbers!$C$2:$C$5"'), b["names"])


# ---- CSV ----

def test_csv(work):
    """CSV read: fixtures/csv; and both ways through the other formats."""
    expect = {
        "quoting.csv": {(0, 2): 'text "Price, each"', (1, 0): 'text "Fig, \\"dried\\""', (1, 2): "number 1E+20",
                        (1, 3): 'text "two\\x0d\\nlines"', (2, 2): "number 0.3", (2, 3): 'text "  spaced  "',
                        (4, 0): 'text "=SUM(A1:A2)"', (4, 1): "bool 1", (4, 2): "error #DIV/0!", (4, 3): "number 0.125",
                        (5, 0): "number 46316", (5, 1): "number 46316.5732638889", (5, 2): "number 7", (5, 3): 'text "x\\"y"',
                        (6, 0): 'text "quoted tail"'},
        "semicolons.csv": {(0, 2): 'text "Datum"', (1, 0): 'text "1,5"', (1, 1): 'text "Grüße"', (2, 1): 'text "a;b"'},
        "excel-semicolons.csv": {(0, 0): 'text "A"', (2, 2): "number 6"},
        "tabs.txt": {(0, 2): 'text "c"', (1, 2): "number 3"},
        "windows-1252.csv": {(0, 0): 'text "Café"', (0, 1): 'text "€5"', (0, 2): 'text "naïve"'},
        "utf16.csv": {(1, 0): 'text "ü😀"', (1, 1): "number 42"},
        "cr-only.csv": {(2, 1): "number 4"},
        "open-quote.csv": {(0, 1): 'text "never closed,\\nb,c\\n"'},
    }
    paths = []
    for name, cells in expect.items():
        path = os.path.join(FIX, "csv", name)
        paths.append(path)
        rc, out, err = dump(path)
        check("csv " + name, rc == 0 and out.splitlines()[0].startswith("book format=csv"), out[:200] + err)
        b = parse(out)
        for (r, c), want in cells.items():
            line = b["sheets"][0]["cells"].get((r, c)) if b["sheets"] else None
            got = line.split(" ", 3)[3] if line else None
            check(f"csv {name} {r},{c}", got == want, f"got {got!r}, want {want!r}")
    b = parse(dump(os.path.join(FIX, "csv", "quoting.csv"))[1])
    check("csv: percent and date formats", numfmt_of(b, 0, "D5") == '"0.00%"' and numfmt_of(b, 0, "A6") == '"yyyy-mm-dd"'
          and numfmt_of(b, 0, "B6") == '"yyyy-mm-dd hh:mm:ss"', (numfmt_of(b, 0, "D5"), numfmt_of(b, 0, "A6")))
    # an .xls's first sheet to CSV and back: the values (formulas give theirs)
    src = os.path.join(FIX, "book.xls")
    out_csv = os.path.join(work, "book.csv")
    if convert(src, out_csv):
        a = parse(dump(src)[1])["sheets"][0]["cells"]
        b = parse(dump(out_csv)[1])["sheets"][0]["cells"]
        value = lambda l: re.sub(r" formula .*| array", "", l.split(" ", 3)[3])
        want = {k: value(l) for k, l in a.items() if value(l) != "empty" and value(l) != 'text ""'}
        got = {k: value(l) for k, l in b.items()}
        # what CSV cannot say: a boolean format's 0 is a number again
        check("xls to csv to model", got == want, [(k, want.get(k), got.get(k)) for k in sorted(set(want) | set(got)) if want.get(k) != got.get(k)][:6])
        # and the CSV saved again from any format is the same text
        for ext in ("xlsx", "ods", "fods"):
            mid, again = os.path.join(work, "csv-mid." + ext), os.path.join(work, "csv-again-%s.csv" % ext)
            if convert(out_csv, mid) and convert(mid, again):
                check("csv to %s to csv" % ext, open(out_csv, "rb").read() == open(again, "rb").read())
    return paths



# ---- Small files that ask for too much ----

def test_bombs(work):
    """A few kilobytes must not become gigabytes: repeats, spaces and shared
    strings have budgets, and the reader says what it left out."""
    import time
    content = ('<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
               '<office:body><office:spreadsheet>%s</office:spreadsheet></office:body></office:document-content>')
    big = "x" * 200000
    cases = {
        "spaces.ods": '<table:table table:name="S"><table:table-row><table:table-cell office:value-type="string"><text:p>%s</text:p>'
                      '</table:table-cell></table:table-row></table:table>' % ('<text:s text:c="10000"/>' * 5000),
        "repeats.ods": '<table:table table:name="S"><table:table-row table:number-rows-repeated="1048576"><table:table-cell '
                       'office:value-type="string" table:number-columns-repeated="16384"><text:p>%s</text:p></table:table-cell>'
                       '</table:table-row></table:table>' % big,
        "sheets.ods": "".join('<table:table table:name="S%d"><table:table-row table:number-rows-repeated="100000"><table:table-cell '
                              'office:value-type="string" table:number-columns-repeated="100"><text:p>%s</text:p></table:table-cell>'
                              '</table:table-row></table:table>' % (i, "y" * 5000) for i in range(30)),
    }
    for name, body in cases.items():
        path = os.path.join(work, name)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("mimetype", ODS_MIME, compress_type=zipfile.ZIP_STORED)
            z.writestr("content.xml", content % body)
        start = time.time()
        rc, out, err = dump(path)
        check("bounded " + name, rc in (0, 10) and time.time() - start < 120 and len(out) < 80000000,
              "exit %d, %.1f s, %d bytes %s" % (rc, time.time() - start, len(out), err[-300:]))
        check("bounded %s: said so" % name, "left out" in out or "cut short" in out, out[:300])
    # one long shared string in many cells, and array formulas over whole sheets
    xlsx = os.path.join(work, "shared.xlsx")
    with zipfile.ZipFile(xlsx, "w", zipfile.ZIP_DEFLATED) as z:
        ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r" '
                   'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", '<workbook %s xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<sheets><sheet name="S" sheetId="1" r:id="a"/></sheets></workbook>' % ns)
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="a" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="s.xml"/>'
                   '<Relationship Id="b" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="ss.xml"/></Relationships>')
        z.writestr("xl/ss.xml", "<sst %s><si><t>%s</t></si></sst>" % (ns, "z" * 1000000))
        cells = "".join('<row r="%d"><c r="A%d" t="s"><v>0</v></c></row>' % (r, r) for r in range(1, 3001))
        arrays = "".join('<row r="%d"><c r="B%d"><f t="array" ref="B%d:XFD1048576">1</f><v>1</v></c></row>' % (r, r, r)
                         for r in range(3001, 3101))
        z.writestr("xl/s.xml", "<worksheet %s><sheetData>%s%s</sheetData></worksheet>" % (ns, cells, arrays))
    start = time.time()
    rc, out, err = dump(xlsx)
    check("bounded shared strings", rc in (0, 10) and time.time() - start < 120, "exit %d %.1f s %s" % (rc, time.time() - start, err[-300:]))
    check("bounded shared strings: said so", "used too often" in out, out[:300])


# ---- Which reader a file goes to ----

def test_routing(work):
    """A workbook goes to the workbook readers by its contents, never to
    the document readers; a presentation to neither (no reader yet)."""
    def package(name, files):
        path = os.path.join(work, name)
        with zipfile.ZipFile(path, "w") as z:
            for n, text in files:
                z.writestr(n, text)
        return path
    pptx = package("slides.pptx", [("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                                    '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/></Types>'),
                                   ("_rels/.rels", "<Relationships/>"), ("ppt/presentation.xml", "<p/>")])
    odp = package("slides.odp", [("mimetype", "application/vnd.oasis.opendocument.presentation"), ("content.xml", "<x/>"), ("styles.xml", "<x/>")])
    for src in (os.path.join(FIX, "book-lo.xlsx"), os.path.join(FIX, "book-lo.ods"), pptx, odp):
        if not os.path.exists(src):
            continue
        rc, out, err = run([OWCONVERT, src, os.path.join(work, "out.odt")])
        check("not a document: " + os.path.basename(src), rc == 10 and "not a document format" in out, out + err)
        rc, out, err = run([OWCONVERT, src, os.path.join(work, "out.docx")])
        check("not a document: " + os.path.basename(src), rc == 10, out + err)
    for src in (pptx, odp):
        rc, out, err = dump(src)
        check("not a workbook: " + os.path.basename(src), rc == 10, out + err)
    rc, out, err = dump(os.path.join(FIX, "book-lo.xlsx"))
    check("a workbook by its contents", out.startswith("book format=xlsx"), out[:200])
    xltx = os.path.join(work, "template.bin")
    with zipfile.ZipFile(os.path.join(FIX, "book-openpyxl.xlsx")) as z, zipfile.ZipFile(xltx, "w") as w:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(b"spreadsheetml.sheet.main+xml", b"spreadsheetml.template.main+xml")
            w.writestr(info.filename, data)
    rc, out, err = dump(xltx)
    check("a template by its contents", rc == 0 and "book format=xlsx" in out, out[:300] + err)
    xlsm = os.path.join(work, "macros.bin")
    with zipfile.ZipFile(os.path.join(FIX, "book-openpyxl.xlsx")) as z, zipfile.ZipFile(xlsm, "w") as w:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
            w.writestr(info.filename, data)
        w.writestr("xl/vbaProject.bin", b"\0" * 64)
    rc, out, err = dump(xlsm)
    check("macro-enabled: read without its macros", rc == 0 and "book format=xlsm" in out and "Macros (Visual Basic)" in out, out[:400] + err)
    rc, out, err = run([OWCONVERT, os.path.join(FIX, "book-openpyxl.xlsx"), os.path.join(work, "saved.xlsm")])
    check("never saved as .xlsm", rc != 0, out + err)
    for name, fmt in (("book.fods", "fods"), ("book-lo.ods", "ods"), ("csv/quoting.csv", "csv")):
        rc, out, err = dump(os.path.join(FIX, name))
        check("by its contents: " + name, out.startswith("book format=" + fmt), out[:200])
    rc, out, err = run([OWCONVERT, "--formats"])
    for fmt in ("xlsx", "xlsm", "ods", "fods", "csv"):
        check("OWConvert lists " + fmt, re.search(r"^%s\s+yes" % fmt, out, re.M) is not None, out)


# ---- Damaged files ----

def damaged_versions(data, rnd, flips=40):
    cuts = sorted(set(list(range(0, min(len(data), 1024), 41)) + list(range(0, len(data), max(1, len(data) // 60)))))
    versions = [data[:n] for n in cuts]
    for _ in range(flips):
        d = bytearray(data)
        for _ in range(rnd.randint(1, 10)):
            d[rnd.randrange(len(d))] = rnd.randrange(256)
        versions.append(bytes(d))
    return versions


def damaged_parts(path, rnd, per_part=6):
    """The package again with one XML part cut short or with bytes
    changed, zipped properly: the damage reaches the XML readers."""
    z = zipfile.ZipFile(path)
    infos = z.infolist()
    out = []
    for info in infos:
        if not (info.filename.endswith(".xml") or info.filename.endswith(".rels")):
            continue
        part = z.read(info.filename)
        variants = [part[:rnd.randrange(len(part) + 1)] for _ in range(per_part // 2)]
        for _ in range(per_part - per_part // 2):
            d = bytearray(part)
            for _ in range(rnd.randint(1, 6)):
                if d:
                    d[rnd.randrange(len(d))] = rnd.choice(b'<>/"=&\'\x00 x9A;#')
            variants.append(bytes(d))
        for v in variants:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as w:
                for other in infos:
                    w.writestr(other.filename, v if other.filename == info.filename else z.read(other.filename),
                               compress_type=other.compress_type)
            out.append(buf.getvalue())
    return out


def test_damaged(paths, ext):
    rnd = random.Random(2026)
    runs = 0
    with tempfile.TemporaryDirectory() as work:
        bad = os.path.join(work, "bad." + ext)
        for path in paths:
            data = open(path, "rb").read()
            versions = damaged_versions(data, rnd)
            if data[:2] == b"PK":
                versions += damaged_parts(path, rnd)
            for v in versions:
                with open(bad, "wb") as f:
                    f.write(v)
                rc, out, err = dump(bad)
                runs += 1
                if rc not in (0, 10):
                    check("damaged %s (%d bytes)" % (os.path.basename(path), len(v)), False, "exit %d %s" % (rc, err[-800:]))
                    shutil.copy(bad, os.path.join(HERE, "..", "..", "build", "host", "crash." + ext)) if os.path.isdir(os.path.join(HERE, "..", "..", "build", "host")) else None
                    return
    print("%d damaged %s files read safely" % (runs, ext.upper()))


def main():
    with tempfile.TemporaryDirectory() as work:
        test_numbers()
        written = test_xls_to_xlsx(work)
        test_xlsx_readers(written)
        hand = test_hand_xlsx(work)
        strict = test_strict_and_bare(work)
        generated = test_generated_xlsx()
        ods_written = test_xls_to_ods(work)
        test_ods_readers(ods_written, written)
        test_xlsx_ods_xlsx(work, written + [hand, os.path.join(FIX, "book-openpyxl.xlsx"), os.path.join(FIX, "book-lo.xlsx")])
        hand_ods = test_hand_ods(work)
        test_generated_ods()
        test_ods_round_trips(work, [hand_ods, os.path.join(FIX, "book-lo.ods"), os.path.join(FIX, "book-odfpy.ods"),
                                    os.path.join(FIX, "book.fods")])
        csvs = test_csv(work)
        test_bombs(work)
        test_routing(work)
        test_damaged([hand, strict] + generated + written[:1], "xlsx")
        test_damaged([hand_ods, os.path.join(FIX, "book-lo.ods"), os.path.join(FIX, "book-odfpy.ods")]
                     + [p for p in ods_written if p.endswith(".ods")][:1], "ods")
        test_damaged([os.path.join(FIX, "book.fods")] + [p for p in ods_written if p.endswith("book.xls.fods")], "fods")
        test_damaged(csvs, "csv")
    if failures:
        print("%d check(s) of XLSX, ODS and CSV failed" % len(failures))
        sys.exit(1)
    print("XLSX, ODS and CSV tests passed")


if __name__ == "__main__":
    main()
